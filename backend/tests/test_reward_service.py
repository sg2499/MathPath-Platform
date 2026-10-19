"""Rewards marketplace (2026-10-18): the catalogue, unlocking with coins,
cancelling, preparing, handing over, rejecting, goals and the ledger."""
import io
import os
import threading
import uuid

import pytest
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base
from app.models import (
    EconomyTransaction,
    Level,
    Module,
    RewardItem,
    RewardOrder,
    Student,
    User,
    UserEconomy,
)
from app.services import reward_service as R
from app.services.economy_service import EconomyService
from app.services.reward_images import RewardImageError, SquareWebp
from app.services.reward_seed import OPENING_ITEMS, SeedRewardItems


@pytest.fixture()
def db(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)  # seed photos land in a scratch uploads/ folder
    engine = create_engine("sqlite://", future=True, connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)()
    session.add(Module(id="m1", module_code="IM", module_name="Intermediate Module"))
    session.add(Level(id="l1", module_id="m1", level_code="IM-L2", level_name="IM Level 2"))
    session.add(User(id="au", full_name="Office Admin", email="a@x.test", password_hash="x", role="ADMIN", is_active=True))
    session.commit()
    yield session
    session.close()


def _student(db, n, coins=0, *, demo=False, active=True, name=None):
    db.add(User(id=f"u{n}", full_name=name or f"Student {n:03d}", email=f"s{n}@x.test", password_hash="x", role="STUDENT", is_active=active))
    db.add(Student(id=f"s{n}", user_id=f"u{n}", student_code=f"MP-ST-{n:03d}", is_active=active, is_demo=demo, current_level_id="l1"))
    db.commit()
    if coins:
        EconomyService.award_xp_and_coins(db, f"u{n}", 0, coins, "DPS_COMPLETION", f"att-{n}")
    return db.get(Student, f"s{n}")


def _item(db, price=800, stock=None, name="Gel Pens", active=True):
    return R.CreateItem(db, Name=name, Price=price, Stock=stock, IsActive=active, UserId="au")


def _balance(db, user_id):
    db.expire_all()
    return db.query(UserEconomy).filter(UserEconomy.user_id == user_id).one().coin_balance


def _ledger(db, user_id):
    return sum(t.amount_coins for t in db.query(EconomyTransaction).filter(EconomyTransaction.user_id == user_id))


# ---------- catalogue ------------------------------------------------------------

def test_opening_catalogue_is_seeded_once_with_photos_and_never_overwrites(db, tmp_path):
    assert SeedRewardItems(db) == 16
    items = R.ListItems(db)
    assert [i.price_coins for i in items] == [p for *_, p in OPENING_ITEMS]
    assert items[0].price_coins == 800 and items[-1].price_coins == 5200
    assert all(i.image_url and i.image_url.startswith("/uploads/rewards/starter-") for i in items)
    assert all((tmp_path / i.image_url.lstrip("/")).exists() for i in items)
    first = items[0]
    R.UpdateItem(db, first.id, Price=900, Name="Edited Kit", IsActive=False)
    assert SeedRewardItems(db) == 0
    again = db.get(RewardItem, first.id)
    assert again.price_coins == 900 and again.name == "Edited Kit" and again.is_active is False
    assert db.query(RewardItem).count() == 16


def test_item_checks(db):
    with pytest.raises(R.RewardError, match="name"):
        R.CreateItem(db, Name="  ", Price=100)
    with pytest.raises(R.RewardError, match="Price"):
        R.CreateItem(db, Name="Pen", Price=0)
    with pytest.raises(R.RewardError, match="whole number"):
        R.CreateItem(db, Name="Pen", Price="lots")
    with pytest.raises(R.RewardError, match="Stock"):
        R.CreateItem(db, Name="Pen", Price=100, Stock=-1)
    item = R.CreateItem(db, Name="  Parker   Pen ", Price="3600", Stock="", Description="  A classic  pen ")
    assert item.name == "Parker Pen" and item.price_coins == 3600 and item.stock is None and item.description == "A classic pen"
    R.UpdateItem(db, item.id, Stock=None)
    assert db.get(RewardItem, item.id).stock is None


def test_reorder_puts_the_given_items_first(db):
    a, _, c = _item(db, 100, name="A"), _item(db, 200, name="B"), _item(db, 300, name="C")
    R.ReorderItems(db, [c.id, a.id])
    assert [i.name for i in R.ListItems(db)] == ["C", "A", "B"]


def test_photos_become_square_webp_and_bad_files_are_refused():
    raw = io.BytesIO()
    Image.new("RGB", (600, 300), (200, 30, 30)).save(raw, "JPEG")
    out = Image.open(io.BytesIO(SquareWebp(raw.getvalue())))
    assert out.format == "WEBP" and out.size == (800, 800)
    with pytest.raises(RewardImageError, match="isn't a photo"):
        SquareWebp(b"not an image at all")
    small = io.BytesIO()
    Image.new("RGB", (100, 100)).save(small, "PNG")
    with pytest.raises(RewardImageError, match="too small"):
        SquareWebp(small.getvalue())
    with pytest.raises(RewardImageError, match="5 MB"):
        SquareWebp(b"x" * (5 * 1024 * 1024 + 1))


# ---------- unlocking --------------------------------------------------------------

def test_unlock_spends_once_and_keeps_the_item_as_it_was(db):
    s = _student(db, 1, 1000)
    item = _item(db, 800, stock=3, name="Writing Kit")
    order, created = R.UnlockReward(db, s, item.id, "key-1")
    assert created and order.status == "UNLOCKED" and order.price_coins == 800
    assert len(order.pickup_code) == 7 and order.pickup_code[3] == "-"
    assert not set(order.pickup_code.replace("-", "")) & set("01OIL")
    assert _balance(db, "u1") == 200 and _ledger(db, "u1") == 200
    assert db.get(RewardItem, item.id).stock == 2
    spend = db.get(EconomyTransaction, order.spend_transaction_id)
    assert spend.transaction_type == "SPEND" and spend.amount_coins == -800 and spend.reference_id == order.id
    econ = db.query(UserEconomy).filter(UserEconomy.user_id == "u1").one()
    assert econ.lifetime_coins_earned == 1000 and econ.current_xp == 0  # spending never lowers these
    # a price change later never touches the order
    R.UpdateItem(db, item.id, Price=1500, Name="Renamed")
    db.refresh(order)
    assert order.price_coins == 800 and order.item_name == "Writing Kit"


def test_the_same_request_twice_unlocks_once(db):
    s = _student(db, 1, 2000)
    item = _item(db, 800)
    first, created1 = R.UnlockReward(db, s, item.id, "tap-1")
    again, created2 = R.UnlockReward(db, s, item.id, "tap-1")
    assert created1 and not created2 and again.id == first.id
    assert _balance(db, "u1") == 1200
    other, created3 = R.UnlockReward(db, s, item.id, "tap-2")  # a new tap is a new unlock
    assert created3 and other.id != first.id and _balance(db, "u1") == 400


def test_not_enough_coins_says_how_many_more(db):
    s = _student(db, 1, 500)
    item = _item(db, 800)
    with pytest.raises(R.RewardError) as caught:
        R.UnlockReward(db, s, item.id, "k")
    assert caught.value.code == "NOT_ENOUGH_COINS" and "300 more coins" in caught.value.message
    assert _balance(db, "u1") == 500 and db.query(RewardOrder).count() == 0


def test_out_of_stock_and_switched_off_items_cannot_be_unlocked(db):
    s = _student(db, 1, 5000)
    empty = _item(db, 100, stock=0, name="Gone")
    off = _item(db, 100, name="Off", active=False)
    with pytest.raises(R.RewardError) as caught:
        R.UnlockReward(db, s, empty.id)
    assert caught.value.code == "OUT_OF_STOCK"
    with pytest.raises(R.RewardError) as caught:
        R.UnlockReward(db, s, off.id)
    assert caught.value.code == "NOT_FOUND"
    last = _item(db, 100, stock=1, name="Last One")
    R.UnlockReward(db, s, last.id)
    with pytest.raises(R.RewardError) as caught:
        R.UnlockReward(db, s, last.id)
    assert caught.value.code == "OUT_OF_STOCK"
    assert _balance(db, "u1") == 4900


def test_exactly_enough_coins_leaves_zero(db):
    s = _student(db, 1, 800)
    R.UnlockReward(db, s, _item(db, 800).id)
    assert _balance(db, "u1") == 0


# ---------- cancelling, preparing, handing over, rejecting -------------------------------

def test_student_cancels_while_unlocked_and_gets_every_coin_back(db):
    s = _student(db, 1, 1000)
    item = _item(db, 800, stock=1)
    order, _ = R.UnlockReward(db, s, item.id)
    back = R.CancelOrder(db, s, order.id)
    assert back.status == "CANCELLED" and back.closed_by_user_id == "u1"
    assert _balance(db, "u1") == 1000 and _ledger(db, "u1") == 1000
    assert db.get(RewardItem, item.id).stock == 1
    refund = db.get(EconomyTransaction, back.refund_transaction_id)
    assert refund.transaction_type == "REFUND" and refund.amount_coins == 800
    assert db.query(UserEconomy).filter(UserEconomy.user_id == "u1").one().lifetime_coins_earned == 1000
    with pytest.raises(R.RewardError) as caught:
        R.CancelOrder(db, s, order.id)  # only once
    assert caught.value.code == "CANNOT_CANCEL"
    assert _balance(db, "u1") == 1000


def test_no_cancelling_once_ready_or_someone_elses(db):
    s1, s2 = _student(db, 1, 1000), _student(db, 2, 1000)
    order, _ = R.UnlockReward(db, s1, _item(db, 800).id)
    with pytest.raises(R.RewardError) as caught:
        R.CancelOrder(db, s2, order.id)
    assert caught.value.code == "NOT_FOUND"
    R.MarkReady(db, [order.id], "au")
    with pytest.raises(R.RewardError) as caught:
        R.CancelOrder(db, s1, order.id)
    assert caught.value.code == "CANNOT_CANCEL" and "at the centre" in caught.value.message
    assert _balance(db, "u1") == 200


def test_prepare_then_hand_over(db):
    s = _student(db, 1, 3000)
    a, _ = R.UnlockReward(db, s, _item(db, 800, name="A").id)
    b, _ = R.UnlockReward(db, s, _item(db, 900, name="B").id)
    moved = R.MarkReady(db, [a.id, b.id, a.id, "nope"], "au")
    assert {o.id for o in moved} == {a.id, b.id}
    assert R.MarkReady(db, [a.id], "au") == []  # already ready: nothing moves
    done = R.MarkCollected(db, a.id, "au")
    assert done.status == "COLLECTED" and done.collected_by_user_id == "au" and done.collected_at is not None
    with pytest.raises(R.RewardError):
        R.MarkCollected(db, a.id, "au")
    with pytest.raises(R.RewardError):
        R.RejectOrder(db, a.id, "too late", "au")
    assert _balance(db, "u1") == 1300


def test_hand_over_straight_from_unlocked(db):
    s = _student(db, 1, 1000)
    order, _ = R.UnlockReward(db, s, _item(db, 800).id)
    done = R.MarkCollected(db, order.id, "au")
    assert done.status == "COLLECTED" and done.ready_at is not None and done.ready_by_user_id == "au"


def test_reject_needs_a_reason_and_refunds(db):
    s = _student(db, 1, 1000)
    item = _item(db, 800, stock=2)
    order, _ = R.UnlockReward(db, s, item.id)
    R.MarkReady(db, [order.id], "au")
    with pytest.raises(R.RewardError, match="reason"):
        R.RejectOrder(db, order.id, "   ", "au")
    done = R.RejectOrder(db, order.id, "  Out of stock  at the shop ", "au")
    assert done.status == "REJECTED" and done.close_reason == "Out of stock at the shop" and done.closed_by_user_id == "au"
    assert _balance(db, "u1") == 1000 and db.get(RewardItem, item.id).stock == 2


def test_find_by_pickup_code_however_it_is_typed(db):
    s = _student(db, 1, 3000, name="Aarav Shah")
    a, _ = R.UnlockReward(db, s, _item(db, 800, name="A").id)
    b, _ = R.UnlockReward(db, s, _item(db, 900, name="B").id)
    typed = a.pickup_code.replace("-", "").lower()
    found = R.FindByCode(db, f" {typed[:3]} {typed[3:]} ")
    assert found["match"]["id"] == a.id and found["student"]["name"] == "Aarav Shah"
    assert [o["id"] for o in found["open"]] == [a.id, b.id]
    assert R.FindByCode(db, "ZZZ-ZZZ") is None


# ---------- the store, goal, history ---------------------------------------------------

def test_store_shows_how_close_every_reward_is(db):
    s = _student(db, 1, 1000)
    _item(db, 800, name="Cheap"), _item(db, 4000, name="Dear"), _item(db, 500, stock=0, name="Empty")
    _item(db, 100, name="Hidden", active=False)
    store = R.StorePayload(db, s)
    assert store["balance"] == 1000 and store["lifetimeEarned"] == 1000 and store["spent"] == 0
    by = {i["name"]: i for i in store["items"]}
    assert set(by) == {"Cheap", "Dear", "Empty"}
    assert by["Cheap"]["canUnlock"] and by["Cheap"]["coinsToGo"] == 0 and by["Cheap"]["progress"] == 1.0
    assert not by["Dear"]["canUnlock"] and by["Dear"]["coinsToGo"] == 3000 and by["Dear"]["progress"] == 0.25
    assert not by["Empty"]["canUnlock"] and by["Empty"]["inStock"] is False  # visible, greyed out
    assert all(i["stock"] is None for i in store["items"])  # students never see stock counts


def test_goal_can_be_set_changed_and_cleared_even_when_out_of_stock(db):
    s = _student(db, 1, 1000)
    a, empty = _item(db, 4000, name="Backpack"), _item(db, 2000, stock=0, name="Football")
    R.SetGoal(db, s, a.id)
    assert R.StorePayload(db, s)["goal"]["name"] == "Backpack"
    R.SetGoal(db, s, empty.id)
    store = R.StorePayload(db, s)
    assert store["goal"]["name"] == "Football" and sum(i["isGoal"] for i in store["items"]) == 1
    R.SetGoal(db, s, None)
    assert R.StorePayload(db, s)["goal"] is None
    off = _item(db, 100, name="Off", active=False)
    with pytest.raises(R.RewardError):
        R.SetGoal(db, s, off.id)


def test_coin_history_in_plain_words(db):
    s = _student(db, 1, 1000)
    EconomyService.award_xp_and_coins(db, "u1", 10, 140, "MOCK_COMPLETION", "m1")
    order, _ = R.UnlockReward(db, s, _item(db, 800, name="Parker Pen").id)
    R.CancelOrder(db, s, order.id)
    entries = R.CoinHistory(db, "u1")["entries"]
    labels = [(e["label"], e["coins"]) for e in entries]
    assert ("Unlocked Parker Pen", -800) in labels and ("Coins returned for Parker Pen", 800) in labels
    assert ("DPS Sheet", 1000) in labels and ("Mock Exam", 140) in labels
    assert sum(e["coins"] for e in entries) == _balance(db, "u1") == 1140


# ---------- admin lists -------------------------------------------------------------

def test_admin_lists_follow_the_show_demo_switch_and_the_shopping_list_skips_demos(db, monkeypatch):
    real, demo = _student(db, 1, 2000, name="Real Student"), _student(db, 2, 2000, demo=True, name="Demo Student")
    item = _item(db, 800, name="Pens")
    R.UnlockReward(db, real, item.id)
    R.UnlockReward(db, demo, item.id)
    assert [o["student"]["name"] for o in R.AdminOrders(db)] == ["Real Student"]
    assert R.ShoppingList(db) == [{"itemId": item.id, "itemName": "Pens", "imageUrl": None, "count": 1, "oldest": R.ShoppingList(db)[0]["oldest"]}]
    coins = R.AdminCoins(db)
    assert [c["name"] for c in coins] == ["Real Student"] and coins[0]["balance"] == 1200 and coins[0]["spent"] == 800 and coins[0]["openRewards"] == 1
    # the page shows "Inactive" for a row without this flag (every row on 19 Oct)
    assert coins[0]["isActive"] is True
    import app.services.demo_students as D
    monkeypatch.setattr(D, "IncludeDemo", lambda Db: True)
    assert {o["student"]["name"] for o in R.AdminOrders(db)} == {"Real Student", "Demo Student"}
    assert R.ShoppingList(db)[0]["count"] == 1  # still real students only


def test_admin_orders_filters_and_who_did_what(db):
    s = _student(db, 1, 5000, name="Aarav Shah")
    a, _ = R.UnlockReward(db, s, _item(db, 800, name="Pens").id)
    b, _ = R.UnlockReward(db, s, _item(db, 900, name="Ball").id)
    R.MarkCollected(db, a.id, "au")
    R.CancelOrder(db, s, b.id)
    rows = {o["id"]: o for o in R.AdminOrders(db)}
    assert rows[a.id]["collectedBy"] == "Office Admin" and rows[b.id]["closedBy"] == "The student"
    assert [o["id"] for o in R.AdminOrders(db, Status="COLLECTED")] == [a.id]
    assert [o["id"] for o in R.AdminOrders(db, Status="OPEN")] == []
    assert [o["id"] for o in R.AdminOrders(db, Search="ball")] == [b.id]
    assert [o["id"] for o in R.AdminOrders(db, Search=a.pickup_code.replace("-", ""))] == [a.id]


def test_inactive_students_open_orders_can_still_be_handled(db):
    s = _student(db, 1, 1000)
    order, _ = R.UnlockReward(db, s, _item(db, 800).id)
    s.is_active = False
    db.commit()
    assert R.RejectOrder(db, order.id, "Left the centre", "au").status == "REJECTED"
    assert _balance(db, "u1") == 1000


def test_every_balance_matches_its_ledger(db):
    s = _student(db, 1, 3000)
    i1, i2 = _item(db, 800), _item(db, 1200)
    o1, _ = R.UnlockReward(db, s, i1.id)
    o2, _ = R.UnlockReward(db, s, i2.id)
    R.CancelOrder(db, s, o1.id)
    R.RejectOrder(db, o2.id, "Broken", "au")
    R.UnlockReward(db, s, i1.id)
    assert R.LedgerMismatches(db) == []
    assert _balance(db, "u1") == 2200


# ---------- races on a real database (Postgres row locks) ---------------------------------

PG_URL = os.getenv("REWARDS_PG_URL")


@pytest.fixture()
def pg(tmp_path, monkeypatch):
    if not PG_URL:
        pytest.skip("set REWARDS_PG_URL to a scratch Postgres database to run the race tests")
    monkeypatch.chdir(tmp_path)
    engine = create_engine(PG_URL, future=True, pool_size=20, max_overflow=5)
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    Maker = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)
    with Maker() as s:
        s.add(Module(id="m1", module_code="IM", module_name="Intermediate Module"))
        s.add(Level(id="l1", module_id="m1", level_code="IM-L2", level_name="IM Level 2"))
        s.add(User(id="au", full_name="Office Admin", email="a@x.test", password_hash="x", role="ADMIN", is_active=True))
        s.commit()
    yield Maker
    engine.dispose()


def _race(Maker, jobs):
    """Run every job at the same moment, each on its own connection."""
    gate = threading.Barrier(len(jobs))
    results = [None] * len(jobs)

    def run(i, job):
        with Maker() as s:
            gate.wait()
            try:
                results[i] = job(s)
            except Exception as e:
                results[i] = e

    threads = [threading.Thread(target=run, args=(i, j)) for i, j in enumerate(jobs)]
    [t.start() for t in threads]
    [t.join() for t in threads]
    return results


def test_race_ten_taps_spend_once(pg):
    with pg() as s:
        st = _student(s, 1, 1000)
        item = _item(s, 800)
        sid, iid = st.id, item.id
    def tap(s):
        order, created = R.UnlockReward(s, s.get(Student, sid), iid, "same-tap")
        return order.id, created

    res = _race(pg, [tap for _ in range(10)])
    assert not [r for r in res if isinstance(r, Exception)], res
    assert sum(1 for r in res if r[1]) == 1 and len({r[0] for r in res}) == 1
    with pg() as s:
        assert _balance(s, "u1") == 200 and s.query(RewardOrder).count() == 1


def test_race_two_tabs_cannot_overspend(pg):
    with pg() as s:
        st = _student(s, 1, 1000)
        sid, iid = st.id, _item(s, 800).id
    res = _race(pg, [lambda s, k=k: R.UnlockReward(s, s.get(Student, sid), iid, f"tab-{k}") for k in range(6)])
    ok = [r for r in res if not isinstance(r, Exception)]
    assert len(ok) == 1 and all(isinstance(r, R.RewardError) and r.code == "NOT_ENOUGH_COINS" for r in res if isinstance(r, Exception))
    with pg() as s:
        assert _balance(s, "u1") == 200 and R.LedgerMismatches(s) == []


def test_race_last_item_goes_to_one_student(pg):
    with pg() as s:
        ids = [_student(s, n, 1000).id for n in range(1, 9)]
        iid = _item(s, 100, stock=1).id
    res = _race(pg, [lambda s, sid=sid: R.UnlockReward(s, s.get(Student, sid), iid, str(uuid.uuid4())) for sid in ids])
    winners = [r for r in res if not isinstance(r, Exception)]
    assert len(winners) == 1 and all(r.code == "OUT_OF_STOCK" for r in res if isinstance(r, Exception))
    with pg() as s:
        assert s.get(RewardItem, iid).stock == 0 and s.query(RewardOrder).count() == 1


def test_race_award_and_spend_at_the_same_moment_lose_nothing(pg):
    with pg() as s:
        st = _student(s, 1, 1000)
        sid, iid = st.id, _item(s, 800).id
    jobs = [lambda s: R.UnlockReward(s, s.get(Student, sid), iid, "u")] + [
        (lambda s, k=k: EconomyService.award_xp_and_coins(s, "u1", 1, 50, "DPS_COMPLETION", f"race-{k}")) for k in range(5)
    ]
    res = _race(pg, jobs)
    assert not [r for r in res if isinstance(r, Exception)], res
    with pg() as s:
        assert _balance(s, "u1") == 1000 - 800 + 5 * 50 and R.LedgerMismatches(s) == []


def test_race_cancel_twice_refunds_once(pg):
    with pg() as s:
        st = _student(s, 1, 1000)
        order, _ = R.UnlockReward(s, st, _item(s, 800).id)
        sid, oid = st.id, order.id
    res = _race(pg, [lambda s: R.CancelOrder(s, s.get(Student, sid), oid) for _ in range(5)])
    assert sum(1 for r in res if not isinstance(r, Exception)) == 1
    with pg() as s:
        assert _balance(s, "u1") == 1000 and R.LedgerMismatches(s) == []


# ---------- Part 3: the admin page's summary, export and filters -----------------

def test_admin_summary_counts_real_students_only(db):
    real, demo = _student(db, 1, 5000), _student(db, 2, 5000, demo=True)
    a, b = _item(db, 800, name="A"), _item(db, 900, name="B", stock=0)
    _item(db, 100, name="Off", active=False)
    o1, _ = R.UnlockReward(db, real, a.id)
    o2, _ = R.UnlockReward(db, real, a.id)
    R.UnlockReward(db, demo, a.id)
    R.MarkReady(db, [o2.id], "au")
    o3, _ = R.UnlockReward(db, real, a.id)
    R.MarkCollected(db, o3.id, "au")
    s = R.AdminSummary(db)
    assert s["toPrepare"] == 1 and s["ready"] == 1 and s["collectedThisMonth"] == 1
    assert s["coinsWithStudents"] == 5000 - 3 * 800
    assert s["items"] == {"active": 2, "off": 1, "outOfStock": 1}


def test_rewards_export_has_every_step_and_no_phone_numbers(db):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.services.reward_export import BuildRewardsWorkbook

    s = _student(db, 1, 3000, name="Aarav Shah")
    a, _ = R.UnlockReward(db, s, _item(db, 800, name="Pens").id)
    R.MarkCollected(db, a.id, "au")
    b, _ = R.UnlockReward(db, s, _item(db, 900, name="Ball").id)
    R.RejectOrder(db, b.id, "Sold out", "au")
    sheet = load_workbook(BytesIO(BuildRewardsWorkbook(R.AdminOrders(db)))).active
    head = [c.value for c in sheet[1]]
    assert head[:8] == ["Unlocked", "Student", "Student ID", "Level", "Reward", "Coins", "Pickup Code", "Status"]
    assert not any("phone" in (h or "").lower() or "mobile" in (h or "").lower() for h in head)
    rows = {r[4]: r for r in sheet.iter_rows(min_row=2, values_only=True)}
    assert rows["Pens"][7] == "Collected" and rows["Pens"][11] == "Office Admin" and rows["Pens"][6] == a.pickup_code
    assert rows["Ball"][7] == "Rejected" and rows["Ball"][14] == "Sold out"


def test_coin_levels_lists_levels_of_active_real_students(db):
    _student(db, 1, 0)
    _student(db, 2, 0, demo=True)
    assert R.CoinLevels(db) == [{"id": "l1", "code": "IM-L2"}]


# ---------- Part 4: notifications ----------------------------------------------------

from app.models import Notification  # noqa: E402


def _notes(db, user_id, type_=None):
    q = db.query(Notification).filter(Notification.recipient_user_id == user_id)
    if type_:
        q = q.filter(Notification.type == type_)
    return q.all()


def test_every_step_of_a_reward_is_told_to_the_student_once(db):
    s = _student(db, 1, 3000)
    item = _item(db, 800, name="Parker Pen")
    order, _ = R.UnlockReward(db, s, item.id, "k1")
    R.UnlockReward(db, s, item.id, "k1")  # a double tap sends nothing more
    unlocked = _notes(db, "u1", "REWARD_UNLOCKED")
    assert len(unlocked) == 1 and order.pickup_code in unlocked[0].message
    assert unlocked[0].category == "REWARD" and unlocked[0].color_variant == "AMBER" and unlocked[0].target_route == "/student/rewards?tab=mine"
    R.MarkReady(db, [order.id], "au")
    R.MarkReady(db, [order.id], "au")
    assert len(_notes(db, "u1", "REWARD_READY")) == 1 and "Parker Pen" in _notes(db, "u1", "REWARD_READY")[0].title
    R.MarkCollected(db, order.id, "au")
    assert len(_notes(db, "u1", "REWARD_COLLECTED")) == 1


def test_rejected_reward_tells_the_student_why_and_that_coins_came_back(db):
    s = _student(db, 1, 1000)
    order, _ = R.UnlockReward(db, s, _item(db, 800, name="Ball").id)
    R.RejectOrder(db, order.id, "Sold out this week", "au")
    note = _notes(db, "u1", "REWARD_REJECTED")[0]
    assert "800 coins are back" in note.message and "Sold out this week" in note.message


def test_admins_hear_about_each_unlock_by_a_real_student_only(db):
    real, demo = _student(db, 1, 2000, name="Aarav Shah"), _student(db, 2, 2000, demo=True)
    item = _item(db, 800, name="Pens")
    order, _ = R.UnlockReward(db, real, item.id)
    R.UnlockReward(db, demo, item.id)
    notes = _notes(db, "au", "REWARD_TO_PREPARE")
    assert len(notes) == 1 and notes[0].title == "Aarav Shah Unlocked Pens" and notes[0].target_route == "/admin/rewards?tab=prepare"
    assert order.pickup_code in notes[0].message


def test_goal_reached_is_told_once_when_coins_cross_the_price(db):
    s = _student(db, 1, 500)
    goal = _item(db, 1000, name="Backpack")
    R.SetGoal(db, s, goal.id)
    EconomyService.award_xp_and_coins(db, "u1", 0, 300, "DPS_COMPLETION", "a1")   # 800: not yet
    assert not _notes(db, "u1", "REWARD_GOAL_REACHED")
    EconomyService.award_xp_and_coins(db, "u1", 0, 250, "DPS_COMPLETION", "a2")   # 1050: crosses
    EconomyService.award_xp_and_coins(db, "u1", 0, 100, "DPS_COMPLETION", "a3")   # already past
    notes = _notes(db, "u1", "REWARD_GOAL_REACHED")
    assert len(notes) == 1 and "Backpack" in notes[0].message and notes[0].target_route == f"/student/rewards?item={goal.id}"
    # unlock it, earn it again: told again
    R.UnlockReward(db, s, goal.id)
    EconomyService.award_xp_and_coins(db, "u1", 0, 1000, "DPS_COMPLETION", "a4")
    assert len(_notes(db, "u1", "REWARD_GOAL_REACHED")) == 2


def test_no_goal_note_for_an_out_of_stock_goal_or_without_a_goal(db):
    s = _student(db, 1, 0)
    goal = _item(db, 500, stock=0, name="Gone")
    EconomyService.award_xp_and_coins(db, "u1", 0, 600, "DPS_COMPLETION", "b1")
    R.SetGoal(db, s, goal.id)
    EconomyService.award_xp_and_coins(db, "u1", 0, 600, "DPS_COMPLETION", "b2")
    assert not _notes(db, "u1", "REWARD_GOAL_REACHED")


def test_a_broken_notification_never_undoes_the_reward(db, monkeypatch):
    import app.services.notification_service as N

    def boom(*a, **k):
        raise RuntimeError("notifications are down")

    monkeypatch.setattr(N, "CreateNotification", boom)
    s = _student(db, 1, 1000)
    order, created = R.UnlockReward(db, s, _item(db, 800).id)
    assert created and db.get(RewardOrder, order.id).status == "UNLOCKED" and _balance(db, "u1") == 200
    assert R.MarkCollected(db, order.id, "au").status == "COLLECTED"


def test_reward_notifications_sit_under_achievements_and_the_right_ones_need_action():
    from app.services.notification_tabs import IsActionType, TabFor

    for Type in ("REWARD_UNLOCKED", "REWARD_READY", "REWARD_COLLECTED", "REWARD_REJECTED", "REWARD_GOAL_REACHED", "REWARD_TO_PREPARE"):
        assert TabFor("REWARD", Type) == "ACHIEVEMENTS"
    assert IsActionType("STUDENT", "REWARD_READY") and IsActionType("STUDENT", "REWARD_GOAL_REACHED")
    assert not IsActionType("STUDENT", "REWARD_COLLECTED") and not IsActionType("STUDENT", "REWARD_UNLOCKED")
    assert IsActionType("SUPER_ADMIN", "REWARD_TO_PREPARE") and not IsActionType("TEACHER", "REWARD_TO_PREPARE")
