"""Payments Phase 2 (2026-10-08): invoices -- duplicate rules, idempotency,
snapshots, cancelling, numbering, filters, PDF, Excel and admin-only access."""

import io
import json

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import create_engine
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.dependencies import get_current_user
from app.models import (
    FeeItem,
    Level,
    Module,
    PaymentAuditLog,
    PaymentBusinessProfile,
    PaymentCentre,
    PaymentInvoice,
    PaymentInvoiceBatch,
    PaymentNumberSequence,
    Student,
    User,
)
from app.services.payments import invoice_pdf, numbering
from app.services.payments import invoices_service as inv
from app.services.payments import setup_service as setup
from app.services.payments.invoice_export import BuildInvoicesWorkbook


def _session():
    engine = create_engine("sqlite://", future=True, connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)()


def _code(excinfo):
    return excinfo.value.detail["code"]


def _world(db, *, numbering_set=True):
    admin = User(id="admin-1", full_name="Admin One", email="admin@example.test", password_hash="x", role="ADMIN", is_active=True)
    db.add(admin)
    module = Module(id="mod-im", module_code="IM", module_name="Intermediate Module")
    db.add(module)
    db.flush()
    level = Level(id="lvl-im3", module_id=module.id, level_code="IM-L3", level_name="Intermediate Level 3")
    db.add(level)
    db.commit()
    setup.EnsurePaymentDefaults(db)
    db.commit()
    rajarhat = db.query(PaymentCentre).filter(PaymentCentre.code == "RAJARHAT").one()
    rajarhat.address = "AA-1, Rajarhat, Kolkata"
    profile = db.get(PaymentBusinessProfile, "default")
    profile.registered_address = "12 Registered Road, Kolkata 700001"
    db.commit()

    students = []
    for sid, name, centre, active in [("s1", "Aarav Sen", rajarhat.id, True), ("s2", "Bina Roy", None, True), ("s3", "Chirag Das", None, False)]:
        user = User(id=f"user-{sid}", full_name=name, email=f"{sid}@example.test", password_hash="x", role="STUDENT", is_active=active)
        db.add(user)
        db.flush()
        db.add(Student(
            id=sid, user_id=user.id, student_code=f"MP-{sid.upper()}", is_active=active, centre_id=centre,
            current_level_id=level.id, current_module_id=module.id, father_name="Parent " + name.split()[0],
            father_mobile="9800000000", present_address="Flat 2, Salt Lake",
        ))
    db.commit()
    monthly = setup.CreateFeeItem(db, Name="Monthly Fee", Amount="1100", BillingType="MONTHLY", Actor=admin)
    bag = setup.CreateFeeItem(db, Name="MathPath Bag", Amount="350", BillingType="ONE_TIME", Actor=admin)
    db.commit()
    if numbering_set:
        numbering.SetStartingNumber(db, Key="INVOICE", NextNumber=1037, Actor=admin)
        db.commit()
    return admin, monthly["feeItemId"], bag["feeItemId"]


def _req(fee_ids, students=("s1", "s2"), month=11, year=2026, **extra):
    return {"feeItemIds": list(fee_ids), "studentIds": list(students), "billingMonth": month, "billingYear": year,
            "invoiceDate": "2026-11-01", **extra}


# --- preview and generate ------------------------------------------------------

def test_preview_shows_what_will_be_created_and_what_is_skipped():
    db = _session()
    admin, monthly, bag = _world(db)
    preview = inv.PreviewInvoices(db, Request=_req([monthly, bag], students=("s1", "s2", "s3")))
    assert preview["numberingReady"] is True and preview["nextNumber"] == "MP-INV-001037"
    assert preview["invoiceCount"] == 4 and preview["skippedCount"] == 2
    assert preview["totalDisplay"] == "₹2,900.00"
    assert preview["dueDate"] == "2026-11-11"  # 10 days after the invoice date
    assert {line["reason"] for line in preview["skipped"]} == {"Student is inactive"}
    # Nothing is written by a preview.
    assert db.query(PaymentInvoice).count() == 0


def test_generate_creates_numbered_invoices_with_the_gst_split_and_audit():
    db = _session()
    admin, monthly, bag = _world(db)
    result = inv.GenerateInvoices(db, Request=_req([monthly, bag]), IdempotencyKey="key-00000001", Actor=admin)
    assert result["invoiceCount"] == 4 and result["replayed"] is False
    assert (result["firstNumber"], result["lastNumber"]) == ("MP-INV-001037", "MP-INV-001040")
    rows = db.query(PaymentInvoice).order_by(PaymentInvoice.invoice_number).all()
    # Sorted by student name, then fee order.
    assert [(r.student_id, r.fee_name) for r in rows] == [("s1", "Monthly Fee"), ("s1", "MathPath Bag"), ("s2", "Monthly Fee"), ("s2", "MathPath Bag")]
    fee = rows[0]
    assert (fee.amount_paise, fee.taxable_paise, fee.cgst_paise, fee.sgst_paise) == (110000, 93220, 8390, 8390)
    assert fee.period_key == "2026-11" and rows[1].period_key is None
    assert fee.description == "Math Path Abacus · IM-L3 · Monthly Fee · November 2026"
    assert fee.status == "PENDING" and fee.due_date.isoformat() == "2026-11-11"
    seq = db.get(PaymentNumberSequence, "INVOICE")
    assert seq.next_number == 1041 and seq.last_issued_number == 1040
    assert db.query(PaymentAuditLog).filter(PaymentAuditLog.entity_type == "INVOICE", PaymentAuditLog.action == "CREATE").count() == 4


def test_a_monthly_fee_is_never_invoiced_twice_for_the_same_month():
    db = _session()
    admin, monthly, bag = _world(db)
    inv.GenerateInvoices(db, Request=_req([monthly]), IdempotencyKey="key-00000001", Actor=admin)
    with pytest.raises(HTTPException) as excinfo:
        inv.GenerateInvoices(db, Request=_req([monthly]), IdempotencyKey="key-00000002", Actor=admin)
    assert _code(excinfo) == "NOTHING_TO_INVOICE"
    preview = inv.PreviewInvoices(db, Request=_req([monthly]))
    assert preview["invoiceCount"] == 0
    assert preview["skipped"][0]["reason"] == "Already has Monthly Fee for November 2026 (MP-INV-001037)"
    # A different month is fine.
    result = inv.GenerateInvoices(db, Request=_req([monthly], month=12), IdempotencyKey="key-00000003", Actor=admin)
    assert result["invoiceCount"] == 2


def test_the_database_itself_refuses_a_second_live_monthly_invoice():
    db = _session()
    admin, monthly, bag = _world(db)
    inv.GenerateInvoices(db, Request=_req([monthly], students=("s1",)), IdempotencyKey="key-00000001", Actor=admin)
    original = db.query(PaymentInvoice).one()
    clone = PaymentInvoice(**{c.name: getattr(original, c.name) for c in PaymentInvoice.__table__.columns if c.name not in ("id", "invoice_number")})
    clone.invoice_number = "MP-INV-999999"
    db.add(clone)
    with pytest.raises(IntegrityError):
        db.flush()
    db.rollback()
    # Once the first is cancelled, a replacement is allowed.
    inv.CancelInvoice(db, InvoiceId=original.id, Reason="wrong month", Actor=admin)
    result = inv.GenerateInvoices(db, Request=_req([monthly], students=("s1",)), IdempotencyKey="key-00000002", Actor=admin)
    assert result["invoiceCount"] == 1 and result["firstNumber"] == "MP-INV-001038"


def test_one_time_items_are_skipped_unless_a_repeat_is_allowed():
    db = _session()
    admin, monthly, bag = _world(db)
    inv.GenerateInvoices(db, Request=_req([bag], month=None, year=None), IdempotencyKey="key-00000001", Actor=admin)
    preview = inv.PreviewInvoices(db, Request=_req([bag], month=None, year=None))
    assert preview["invoiceCount"] == 0 and preview["skipped"][0]["reason"].startswith("Already has MathPath Bag (MP-INV-")
    result = inv.GenerateInvoices(db, Request=_req([bag], month=None, year=None, allowRepeatOneTime=True), IdempotencyKey="key-00000002", Actor=admin)
    assert result["invoiceCount"] == 2


def test_the_same_generate_press_never_creates_a_second_set():
    db = _session()
    admin, monthly, bag = _world(db)
    first = inv.GenerateInvoices(db, Request=_req([monthly, bag]), IdempotencyKey="press-abc-123", Actor=admin)
    again = inv.GenerateInvoices(db, Request=_req([monthly, bag]), IdempotencyKey="press-abc-123", Actor=admin)
    assert again["replayed"] is True and again["batchId"] == first["batchId"]
    assert (again["firstNumber"], again["lastNumber"]) == (first["firstNumber"], first["lastNumber"])
    assert db.query(PaymentInvoice).count() == 4 and db.query(PaymentInvoiceBatch).count() == 1


@pytest.mark.parametrize(
    "change,code",
    [
        ({"feeItemIds": []}, "NO_FEE_ITEMS_SELECTED"),
        ({"studentIds": []}, "NO_STUDENTS_SELECTED"),
        ({"billingMonth": None}, "BILLING_PERIOD_REQUIRED"),
        ({"billingMonth": 13}, "BILLING_PERIOD_REQUIRED"),
        ({"dueDate": "2026-10-01"}, "DUE_BEFORE_INVOICE"),
        ({"invoiceDate": "01/11/2026"}, "DATE_INVALID"),
        ({"invoiceDate": "2019-01-01"}, "DATE_INVALID"),
        ({"studentIds": ["nobody"]}, "STUDENT_NOT_FOUND"),
        ({"feeItemIds": ["nothing"]}, "FEE_ITEM_NOT_FOUND"),
    ],
)
def test_bad_requests_are_refused_with_a_clear_reason(change, code):
    db = _session()
    admin, monthly, bag = _world(db)
    request = {**_req([monthly]), **change}
    with pytest.raises(HTTPException) as excinfo:
        inv.PreviewInvoices(db, Request=request)
    assert _code(excinfo) == code


def test_switched_off_fee_items_cannot_be_invoiced():
    db = _session()
    admin, monthly, bag = _world(db)
    setup.SetFeeItemActive(db, FeeItemId=bag, IsActive=False, Reason="sold out", Actor=admin)
    db.commit()
    with pytest.raises(HTTPException) as excinfo:
        inv.PreviewInvoices(db, Request=_req([bag]))
    assert _code(excinfo) == "FEE_ITEM_INACTIVE"


def test_bad_idempotency_keys_are_refused():
    db = _session()
    admin, monthly, bag = _world(db)
    for key in ["", "short", "has space in it", "x" * 81]:
        with pytest.raises(HTTPException) as excinfo:
            inv.GenerateInvoices(db, Request=_req([monthly]), IdempotencyKey=key, Actor=admin)
        assert _code(excinfo) == "IDEMPOTENCY_KEY_INVALID"


def test_nothing_is_invoiced_before_the_starting_number_is_set():
    db = _session()
    admin, monthly, bag = _world(db, numbering_set=False)
    preview = inv.PreviewInvoices(db, Request=_req([monthly]))
    assert preview["numberingReady"] is False and preview["nextNumber"] is None
    with pytest.raises(HTTPException) as excinfo:
        inv.GenerateInvoices(db, Request=_req([monthly]), IdempotencyKey="key-00000001", Actor=admin)
    assert _code(excinfo) == "PAYMENT_NUMBERING_NOT_SET"
    assert db.query(PaymentInvoice).count() == 0 and db.query(PaymentInvoiceBatch).count() == 0


# --- snapshot ------------------------------------------------------------------

def test_an_issued_invoice_never_changes_when_prices_students_or_business_change():
    db = _session()
    admin, monthly, bag = _world(db)
    inv.GenerateInvoices(db, Request=_req([monthly], students=("s1", "s2")), IdempotencyKey="key-00000001", Actor=admin)
    s1 = db.query(PaymentInvoice).filter(PaymentInvoice.student_id == "s1").one()
    s2 = db.query(PaymentInvoice).filter(PaymentInvoice.student_id == "s2").one()
    before = inv.GetInvoice(db, s1.id)
    assert before["snapshot"]["centre"]["name"] == "Rajarhat"
    assert before["snapshot"]["student"]["parentName"] == "Parent Aarav"
    assert before["snapshot"]["business"]["registeredAddress"] == "12 Registered Road, Kolkata 700001"
    # No centre: every centre with an address is listed.
    assert [c["name"] for c in inv.GetInvoice(db, s2.id)["snapshot"]["centre"]["all"]] == ["Rajarhat", "Laketown"]

    setup.UpdateFeeItem(db, FeeItemId=monthly, Amount="1500", Actor=admin)
    student = db.get(Student, "s1")
    student.father_name = "Someone Else"
    db.get(User, "user-s1").full_name = "Renamed"
    setup.UpdateBusinessProfile(db, Fields={"registeredAddress": "New address"}, Actor=admin)
    db.commit()
    after = inv.GetInvoice(db, s1.id)
    assert after["snapshot"] == before["snapshot"]
    assert after["amountDisplay"] == "₹1,100.00" and after["studentName"] == "Aarav Sen"


# --- cancel --------------------------------------------------------------------

def test_cancel_needs_a_reason_and_keeps_the_invoice():
    db = _session()
    admin, monthly, bag = _world(db)
    inv.GenerateInvoices(db, Request=_req([bag], students=("s1",), month=None, year=None), IdempotencyKey="key-00000001", Actor=admin)
    row = db.query(PaymentInvoice).one()
    with pytest.raises(HTTPException) as excinfo:
        inv.CancelInvoice(db, InvoiceId=row.id, Reason="   ", Actor=admin)
    assert _code(excinfo) == "REASON_REQUIRED"
    result = inv.CancelInvoice(db, InvoiceId=row.id, Reason="  ordered by mistake ", Actor=admin)
    assert result["status"] == "CANCELLED" and result["cancelReason"] == "ordered by mistake"
    assert result["cancelledByName"] == "Admin One" and result["balancePaise"] == 0
    with pytest.raises(HTTPException) as excinfo:
        inv.CancelInvoice(db, InvoiceId=row.id, Reason="again", Actor=admin)
    assert _code(excinfo) == "INVOICE_ALREADY_CANCELLED"
    entry = db.query(PaymentAuditLog).filter(PaymentAuditLog.action == "CANCEL").one()
    assert entry.reason == "ordered by mistake" and json.loads(entry.before_json)["status"] == "PENDING"
    # The number is not given back.
    assert db.get(PaymentNumberSequence, "INVOICE").next_number == 1038


# 2026-10-08 (Phase 3): an invoice with payments can now be cancelled; what
# was paid on it moves to the student's advance. See
# test_payments_phase3_receipts.py::test_cancelling_a_paid_invoice_moves_the_money_to_advance.


# --- list and filters ----------------------------------------------------------

def test_list_filters_and_totals_leave_out_cancelled():
    db = _session()
    admin, monthly, bag = _world(db)
    inv.GenerateInvoices(db, Request=_req([monthly, bag]), IdempotencyKey="key-00000001", Actor=admin)
    bag_s2 = db.query(PaymentInvoice).filter(PaymentInvoice.student_id == "s2", PaymentInvoice.fee_name == "MathPath Bag").one()
    inv.CancelInvoice(db, InvoiceId=bag_s2.id, Reason="not needed", Actor=admin)

    everything = inv.ListInvoices(db, Filters={})
    assert everything["totalCount"] == 4
    assert everything["totals"]["amountDisplay"] == "₹2,550.00"  # 1100 + 1100 + 350
    assert inv.ListInvoices(db, Filters={"status": "CANCELLED"})["totalCount"] == 1
    assert inv.ListInvoices(db, Filters={"status": "UNPAID"})["totalCount"] == 3
    assert inv.ListInvoices(db, Filters={"period": "2026-11"})["totalCount"] == 2
    assert inv.ListInvoices(db, Filters={"feeItemId": bag})["totalCount"] == 2
    assert inv.ListInvoices(db, Filters={"search": "bina"})["totalCount"] == 2
    assert inv.ListInvoices(db, Filters={"search": "001037"})["totalCount"] == 1
    assert inv.ListInvoices(db, Filters={"centreId": "NONE"})["totalCount"] == 2
    assert inv.ListInvoices(db, Filters={"dateFrom": "2026-11-02"})["totalCount"] == 0
    # Overdue: unpaid and past the due date.
    assert inv.ListInvoices(db, Filters={"status": "OVERDUE"})["totalCount"] == (3 if inv.TodayInIndia().isoformat() > "2026-11-11" else 0)
    page = inv.ListInvoices(db, Filters={}, Page=2, PageSize=3)
    assert len(page["invoices"]) == 1 and page["invoices"][0]["createdByName"] == "Admin One"


def test_student_options_carry_the_filter_fields():
    db = _session()
    _world(db)
    options = {row["studentId"]: row for row in inv.ListStudentsForInvoicing(db)}
    assert options["s1"]["centreName"] == "Rajarhat" and options["s1"]["levelCode"] == "IM-L3" and options["s1"]["moduleCode"] == "IM"
    assert options["s3"]["isActive"] is False


# --- PDF and Excel -------------------------------------------------------------

def _pdf_text(content):
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)


def test_invoice_pdf_has_every_part_of_the_tax_invoice():
    db = _session()
    admin, monthly, bag = _world(db)
    inv.GenerateInvoices(db, Request=_req([monthly], students=("s1",)), IdempotencyKey="key-00000001", Actor=admin)
    row = db.query(PaymentInvoice).one()
    content = invoice_pdf.RenderInvoicesPdf([row])
    assert content.startswith(b"%PDF")
    text = _pdf_text(content)
    for expected in [
        "TAX INVOICE", "MP-INV-001037", "BGM Enterprise", "19AALPG9427A1ZQ", "12 Registered Road", "Rajarhat",
        "Aarav Sen", "Parent Aarav", "IM-L3", "Monthly Fee", "November 2026", "932.20", "83.90", "1,100.00",
        "Rupees One Thousand One Hundred Only", "Authorised Signatory", "computer-generated", "PENDING",
    ]:
        assert expected in text, expected


def test_bulk_pdf_is_one_page_per_invoice_and_marks_cancelled():
    db = _session()
    admin, monthly, bag = _world(db)
    inv.GenerateInvoices(db, Request=_req([monthly, bag]), IdempotencyKey="key-00000001", Actor=admin)
    first = db.query(PaymentInvoice).order_by(PaymentInvoice.invoice_number).first()
    inv.CancelInvoice(db, InvoiceId=first.id, Reason="duplicate", Actor=admin)
    rows = inv.InvoicesForPdf(db, Filters={})
    reader = PdfReader(io.BytesIO(invoice_pdf.RenderInvoicesPdf(rows)))
    assert len(reader.pages) == 4
    assert "CANCELLED" in reader.pages[0].extract_text() and "duplicate" in reader.pages[0].extract_text()
    assert "CANCELLED" not in reader.pages[1].extract_text()


def test_pdf_limit_and_empty_selection():
    db = _session()
    _world(db)
    with pytest.raises(HTTPException) as excinfo:
        inv.InvoicesForPdf(db, InvoiceIds=[f"id-{n}" for n in range(inv.MAX_INVOICES_PER_PDF + 1)])
    assert _code(excinfo) == "TOO_MANY_INVOICES"
    with pytest.raises(HTTPException) as excinfo:
        inv.InvoicesForPdf(db, Filters={})
    assert _code(excinfo) == "INVOICE_NOT_FOUND"


def test_excel_export_has_numbers_and_a_total_without_cancelled():
    db = _session()
    admin, monthly, bag = _world(db)
    inv.GenerateInvoices(db, Request=_req([monthly, bag]), IdempotencyKey="key-00000001", Actor=admin)
    first = db.query(PaymentInvoice).order_by(PaymentInvoice.invoice_number).first()
    inv.CancelInvoice(db, InvoiceId=first.id, Reason="duplicate", Actor=admin)
    book = load_workbook(io.BytesIO(BuildInvoicesWorkbook(inv.InvoicesForExport(db, Filters={}))))
    sheet = book["Invoices"]
    header = [cell.value for cell in sheet[1]]
    amount_col = header.index("Amount (Rs)") + 1
    assert sheet.cell(row=2, column=amount_col).value == 1100.0
    last = sheet.max_row
    assert sheet.cell(row=last, column=1).value == "Total (excluding cancelled)"
    assert sheet.cell(row=last, column=amount_col).value == 1800.0  # 350 + 1100 + 350


# --- API -----------------------------------------------------------------------

def _api(db, user):
    from app.main import app

    app.dependency_overrides[get_current_user] = lambda: user
    app.dependency_overrides[get_db] = lambda: db
    return TestClient(app), app


def test_invoice_api_end_to_end():
    db = _session()
    admin, monthly, bag = _world(db)
    client, app = _api(db, admin)
    try:
        body = _req([monthly])
        preview = client.post("/api/admin/payments/invoices/preview", json=body)
        assert preview.status_code == 200 and preview.json()["invoiceCount"] == 2
        made = client.post("/api/admin/payments/invoices/generate", json={**body, "idempotencyKey": "api-key-0001"})
        assert made.status_code == 200 and made.json()["invoiceCount"] == 2
        listed = client.get("/api/admin/payments/invoices", params={"period": "2026-11"}).json()
        invoice_id = listed["invoices"][0]["invoiceId"]
        assert client.get(f"/api/admin/payments/invoices/{invoice_id}").json()["snapshot"]["business"]["legalName"] == "BGM Enterprise"
        pdf = client.get(f"/api/admin/payments/invoices/{invoice_id}/pdf")
        assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF")
        assert "MP-INV-00103" in pdf.headers["content-disposition"]
        bulk = client.post("/api/admin/payments/invoices/pdf", json={"filters": {"period": "2026-11"}})
        assert bulk.status_code == 200 and len(PdfReader(io.BytesIO(bulk.content)).pages) == 2
        xlsx = client.get("/api/admin/payments/invoices/export", params={"status": "UNPAID"})
        assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
        options = client.get("/api/admin/payments/invoices/student-options")
        assert options.status_code == 200 and len(options.json()["students"]) == 3
        cancel = client.post(f"/api/admin/payments/invoices/{invoice_id}/cancel", json={"reason": "test"})
        assert cancel.status_code == 200 and cancel.json()["status"] == "CANCELLED"
        assert client.post(f"/api/admin/payments/invoices/{invoice_id}/cancel", json={"reason": "test"}).status_code == 409
        assert client.get("/api/admin/payments/invoices/nope").status_code == 404
    finally:
        app.dependency_overrides.clear()
