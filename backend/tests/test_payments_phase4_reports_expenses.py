"""Payments Phase 4 (2026-10-08): Overview, Collections and Dues reports,
expenses and expense categories."""

import io
from datetime import date

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.database import get_db
from app.dependencies import get_current_user
from app.models import Expense, ExpenseCategory, PaymentAuditLog, PaymentInvoice, PaymentNumberSequence, User
from app.services.payments import expenses_service as exp
from app.services.payments import invoice_pdf
from app.services.payments import invoices_service as inv
from app.services.payments import receipts_service as rec
from app.services.payments import reports_service as rep
from app.services.payments.invoice_export import BuildCollectionsWorkbook, BuildDuesWorkbook, BuildExpensesWorkbook
from tests.test_payments_phase2_invoices import _code, _req
from tests.test_payments_phase3_receipts import _invoice, _pay, _setup

TODAY = date(2026, 11, 20)


@pytest.fixture(autouse=True)
def _fixed_today(monkeypatch):
    monkeypatch.setattr(inv, "TodayInIndia", lambda: TODAY)


def _world_with_payments():
    db, admin, monthly, bag = _setup()
    db.add(User(id="admin-2", full_name="Front Desk", email="fd@example.test", password_hash="x", role="ADMIN", is_active=True))
    db.commit()
    fee1, bag1 = _invoice(db, "s1", "Monthly Fee"), _invoice(db, "s1", "MathPath Bag")
    fee2 = _invoice(db, "s2", "Monthly Fee")
    # 20 Nov: s1 pays the fee (cash 500 + UPI 600), received by admin-1.
    _pay(db, admin, allocations=[(fee1.id, 110000, 0)], methods=[("CASH", 50000, None), ("UPI", 60000, "U1")], paymentDate="2026-11-20")
    # 19 Nov: s2 pays 400 of the fee by UPI, received by Front Desk.
    _pay(db, admin, student="s2", allocations=[(fee2.id, 40000, 0)], methods=[("UPI", 40000, "U2")], paymentDate="2026-11-19", receivedByUserId="admin-2")
    # 19 Nov: s1 pays the bag with a discount, plus 1,000 advance, then that payment is cancelled.
    cancelled = _pay(db, admin, allocations=[(bag1.id, 30000, 5000)], methods=[("CASH", 130000, None)], paymentDate="2026-11-19", discountReason="promo", keepAdvance=True)
    rec.CancelPayment(db, PaymentId=cancelled["paymentId"], Reason="entered twice", Actor=admin)
    # 2 Oct: an older payment, outside November.
    inv.GenerateInvoices(db, Request=_req([monthly], students=("s2",), month=9, invoiceDate="2026-09-01"), IdempotencyKey="gen-old-0001", Actor=admin)
    old = db.query(PaymentInvoice).filter(PaymentInvoice.period_key == "2026-09").one()
    _pay(db, admin, student="s2", allocations=[(old.id, 20000, 0)], methods=[("CHEQUE", 20000, "CHQ-1")], paymentDate="2026-10-02")
    return db, admin, monthly, bag


# --- collections ---------------------------------------------------------------

def test_collections_by_method_staff_and_day():
    db, admin, monthly, bag = _world_with_payments()
    report = rep.CollectionsReport(db, Filters={"dateFrom": "2026-11-01", "dateTo": "2026-11-30"})
    assert report["total"]["display"] == "₹1,500.00"  # 1100 + 400; the cancelled one is left out
    assert report["paymentCount"] == 2
    assert {row["methodLabel"]: row["display"] for row in report["byMethod"]} == {"UPI": "₹1,000.00", "Cash": "₹500.00"}
    assert {row["name"]: row["display"] for row in report["byStaff"]} == {"Admin One": "₹1,100.00", "Front Desk": "₹400.00"}
    assert [(row["date"], row["display"]) for row in report["byDay"]] == [("2026-11-19", "₹400.00"), ("2026-11-20", "₹1,100.00")]
    assert [row["paymentDate"] for row in report["payments"]] == ["2026-11-20", "2026-11-19"]  # newest first


def test_collections_filters_count_only_matching_money():
    db, admin, monthly, bag = _world_with_payments()
    cash = rep.CollectionsReport(db, Filters={"dateFrom": "2026-11-01", "dateTo": "2026-11-30", "method": "CASH"})
    assert cash["total"]["display"] == "₹500.00" and cash["payments"][0]["inFilterDisplay"] == "₹500.00"
    staff = rep.CollectionsReport(db, Filters={"dateFrom": "2026-11-01", "dateTo": "2026-11-30", "receivedBy": "admin-2"})
    assert staff["total"]["display"] == "₹400.00"
    centre = rep.CollectionsReport(db, Filters={"dateFrom": "2026-11-01", "dateTo": "2026-11-30", "centreId": "NONE"})
    assert centre["total"]["display"] == "₹400.00"  # s2 has no centre; s1 is Rajarhat
    today = rep.CollectionsReport(db, Filters={})
    assert today["dateFrom"] == today["dateTo"] == "2026-11-20" and today["total"]["display"] == "₹1,100.00"
    october = rep.CollectionsReport(db, Filters={"dateFrom": "2026-10-01", "dateTo": "2026-10-31"})
    assert october["total"]["display"] == "₹200.00"


@pytest.mark.parametrize("filters,code", [({"dateFrom": "2026-11-10", "dateTo": "2026-11-01"}, "DATE_RANGE_INVALID"), ({"dateFrom": "2020-01-01", "dateTo": "2026-11-01"}, "DATE_RANGE_TOO_LONG"), ({"dateFrom": "x"}, "DATE_INVALID")])
def test_collections_bad_ranges(filters, code):
    db, admin, monthly, bag = _setup()
    with pytest.raises(HTTPException) as excinfo:
        rep.CollectionsReport(db, Filters=filters)
    assert _code(excinfo) == code


def test_collections_excel_and_day_close_pdf():
    db, admin, monthly, bag = _world_with_payments()
    report, lines = rep.CollectionLinesForExport(db, Filters={"dateFrom": "2026-11-01", "dateTo": "2026-11-30"})
    book = load_workbook(io.BytesIO(BuildCollectionsWorkbook(report, lines)))
    assert book.sheetnames == ["Summary", "By Staff", "Daily", "Payments"]
    assert [tuple(r) for r in book["Summary"].iter_rows(values_only=True)][-1] == ("Total", 1500.0)
    assert book["Payments"].cell(row=book["Payments"].max_row, column=8).value == 1500.0
    day = rep.CollectionsReport(db, Filters={"dateFrom": "2026-11-20", "dateTo": "2026-11-20"})
    text = PdfReader(io.BytesIO(invoice_pdf.RenderCollectionSummaryPdf(inv._BusinessSnapshot(db), day))).pages[0].extract_text()
    for expected in ["DAY CLOSE", "20 Nov 2026", "Cash", "UPI", "1,100.00", "Admin One", "MP-MRCPT-631", "Aarav Sen", "Cash counted by"]:
        assert expected in text, expected


# --- dues ----------------------------------------------------------------------

def test_dues_ages_and_totals_match_every_account():
    db, admin, monthly, bag = _world_with_payments()
    report = rep.DuesReport(db, Filters={})
    # s1: bag 350 (due 11 Nov, 9 days overdue). s2: Nov fee 700 left + bag 350 (9 days), Sept fee 900 left (due 11 Sep, 70 days).
    assert report["total"]["display"] == "₹2,300.00"
    assert {row["bucket"]: row["display"] for row in report["buckets"]} == {"NOT_DUE": "₹0.00", "D0_30": "₹1,400.00", "D31_60": "₹0.00", "D61_90": "₹900.00", "D90_PLUS": "₹0.00"}
    assert report["overdue"]["display"] == "₹2,300.00" and report["studentCount"] == 2 and report["invoiceCount"] == 4
    first = report["students"][0]
    assert first["studentName"] == "Bina Roy" and first["dueDisplay"] == "₹1,950.00" and first["maxDaysOverdue"] == 70 and first["bucket"] == "D61_90"
    assert first["oldestDueDate"] == "2026-09-11"
    # The report agrees with each student's own account to the paisa.
    for row in report["students"]:
        assert row["duePaise"] == rec.StudentAccount(db, row["studentId"])["totals"]["duePaise"]


def test_dues_filters_and_sorting():
    db, admin, monthly, bag = _world_with_payments()
    assert rep.DuesReport(db, Filters={"bucket": "D61_90"})["total"]["display"] == "₹900.00"
    assert rep.DuesReport(db, Filters={"centreId": "NONE"})["studentCount"] == 1
    assert rep.DuesReport(db, Filters={"search": "parent aarav"})["students"][0]["studentName"] == "Aarav Sen"
    assert rep.DuesReport(db, Filters={"feeItemId": bag})["total"]["display"] == "₹700.00"
    assert rep.DuesReport(db, Filters={"levelCode": "IM-L3"})["studentCount"] == 2
    assert rep.DuesReport(db, Filters={"levelCode": "PM-L1"})["studentCount"] == 0
    assert [row["studentName"] for row in rep.DuesReport(db, Filters={"sort": "name"})["students"]] == ["Aarav Sen", "Bina Roy"]
    book = load_workbook(io.BytesIO(BuildDuesWorkbook(rep.DuesReport(db, Filters={}))))
    assert book.sheetnames == ["Students", "Invoices", "By Age"]
    assert book["Students"].cell(row=book["Students"].max_row, column=13).value == 2300.0


def test_bucket_boundaries():
    assert [rep.BucketFor(days) for days in (-5, 0, 1, 30, 31, 60, 61, 90, 91, 400)] == ["NOT_DUE", "NOT_DUE", "D0_30", "D0_30", "D31_60", "D31_60", "D61_90", "D61_90", "D90_PLUS", "D90_PLUS"]


# --- overview ------------------------------------------------------------------

def test_overview_adds_up():
    db, admin, monthly, bag = _world_with_payments()
    exp.CreateExpense(db, Request={"expenseDate": "2026-11-05", "categoryId": db.query(ExpenseCategory).filter_by(name="Rent").one().id, "item": "November rent", "methods": [{"method": "UPI", "amountPaise": 30000}]}, IdempotencyKey="exp-00000001", Actor=admin)
    view = rep.Overview(db)
    assert view["todayCollected"]["display"] == "₹1,100.00" and view["todayPaymentCount"] == 1
    assert view["thisMonthCollected"]["display"] == "₹1,500.00" and view["lastMonthCollected"]["display"] == "₹200.00"
    assert view["thisMonthSpent"]["display"] == "₹300.00" and view["thisMonthNet"]["display"] == "₹1,200.00"
    assert view["due"]["display"] == "₹2,300.00" and view["studentsWithDues"] == 2
    assert [row["month"] for row in view["series"]] == ["2026-06", "2026-07", "2026-08", "2026-09", "2026-10", "2026-11"]
    assert view["series"][-1]["collected"]["display"] == "₹1,500.00" and view["series"][-1]["spent"]["display"] == "₹300.00"
    assert view["advanceHeld"]["paise"] == rec.StudentAdvanceBalance(db, "s1") + rec.StudentAdvanceBalance(db, "s2") == 0
    assert len(view["recentPayments"]) == 4


# --- expenses ------------------------------------------------------------------

def _category(db, name="Rent"):
    return db.query(ExpenseCategory).filter_by(name=name).one()


def test_expense_categories_start_list_and_rules():
    db, admin, monthly, bag = _setup()
    names = [row["name"] for row in exp.ListCategories(db)]
    assert names == exp.DEFAULT_CATEGORIES
    added = exp.CreateCategory(db, Name="  Exam   Fees ", Actor=admin)
    assert added["name"] == "Exam Fees"
    with pytest.raises(HTTPException) as excinfo:
        exp.CreateCategory(db, Name="rent", Actor=admin)
    assert _code(excinfo) == "CATEGORY_EXISTS"
    with pytest.raises(HTTPException) as excinfo:
        exp.UpdateCategory(db, CategoryId=added["categoryId"], IsActive=False, Actor=admin)
    assert _code(excinfo) == "REASON_REQUIRED"
    off = exp.UpdateCategory(db, CategoryId=added["categoryId"], IsActive=False, Reason="not needed", Actor=admin)
    assert off["isActive"] is False
    assert db.query(PaymentAuditLog).filter(PaymentAuditLog.entity_type == "EXPENSE_CATEGORY").count() == 2


def test_expense_create_numbering_and_rename_keeps_old_name():
    db, admin, monthly, bag = _setup()
    rent = _category(db)
    request = {"expenseDate": "2026-11-05", "categoryId": rent.id, "item": "November rent", "vendor": "Landlord", "billNumber": "R-11",
               "methods": [{"method": "CASH", "amountPaise": 500000}, {"method": "UPI", "amountPaise": 700000, "reference": "UTR9"}]}
    first = exp.CreateExpense(db, Request=request, IdempotencyKey="exp-00000001", Actor=admin)
    assert first["expenseNumber"] == "MP-EXP-0001" and first["amountDisplay"] == "₹12,000.00" and first["categoryName"] == "Rent"
    again = exp.CreateExpense(db, Request=request, IdempotencyKey="exp-00000001", Actor=admin)
    assert again["replayed"] is True and db.query(Expense).count() == 1
    exp.UpdateCategory(db, CategoryId=rent.id, Name="Centre Rent", Actor=admin)
    assert exp.GetExpense(db, first["expenseId"])["categoryName"] == "Rent"
    second = exp.CreateExpense(db, Request={**request, "categoryId": rent.id}, IdempotencyKey="exp-00000002", Actor=admin)
    assert second["expenseNumber"] == "MP-EXP-0002" and second["categoryName"] == "Centre Rent"
    assert db.get(PaymentNumberSequence, "EXPENSE").is_configured is True


@pytest.mark.parametrize(
    "change,code",
    [
        ({"categoryId": None}, "CATEGORY_REQUIRED"),
        ({"categoryId": "nope"}, "CATEGORY_NOT_FOUND"),
        ({"item": "  "}, "ITEM_REQUIRED"),
        ({"expenseDate": "2026-12-01"}, "DATE_IN_FUTURE"),
        ({"methods": []}, "METHOD_REQUIRED"),
        ({"methods": [{"method": "CASH", "amountPaise": 0}]}, "AMOUNT_ZERO"),
        ({"methods": [{"method": "RAZORPAY", "amountPaise": 100}]}, "METHOD_INVALID"),
        ({"centreId": "nope"}, "CENTRE_NOT_FOUND"),
    ],
)
def test_bad_expenses_are_refused(change, code):
    db, admin, monthly, bag = _setup()
    request = {"expenseDate": "2026-11-05", "categoryId": _category(db).id, "item": "Rent", "methods": [{"method": "CASH", "amountPaise": 1000}], **change}
    with pytest.raises(HTTPException) as excinfo:
        exp.CreateExpense(db, Request=request, IdempotencyKey="exp-bad-0001", Actor=admin)
    assert _code(excinfo) == code


def test_switched_off_category_cannot_be_used_for_new_expenses():
    db, admin, monthly, bag = _setup()
    travel = _category(db, "Travel")
    exp.UpdateCategory(db, CategoryId=travel.id, IsActive=False, Reason="merge", Actor=admin)
    with pytest.raises(HTTPException) as excinfo:
        exp.CreateExpense(db, Request={"categoryId": travel.id, "item": "Cab", "methods": [{"method": "CASH", "amountPaise": 100}]}, IdempotencyKey="exp-00000009", Actor=admin)
    assert _code(excinfo) == "CATEGORY_INACTIVE"


def test_expense_edit_and_cancel_with_history():
    db, admin, monthly, bag = _setup()
    made = exp.CreateExpense(db, Request={"expenseDate": "2026-11-05", "categoryId": _category(db, "Printing").id, "item": "Worksheets", "methods": [{"method": "CASH", "amountPaise": 80000}]}, IdempotencyKey="exp-00000001", Actor=admin)
    with pytest.raises(HTTPException) as excinfo:
        exp.EditExpense(db, ExpenseId=made["expenseId"], Request={}, Reason="", Actor=admin)
    assert _code(excinfo) == "REASON_REQUIRED"
    edited = exp.EditExpense(db, ExpenseId=made["expenseId"], Request={"expenseDate": "2026-11-06", "categoryId": _category(db, "Printing").id, "item": "Worksheets (IM-3)", "methods": [{"method": "UPI", "amountPaise": 85000, "reference": "U5"}]}, Reason="bill corrected", Actor=admin)
    assert edited["amountDisplay"] == "₹850.00" and edited["expenseNumber"] == "MP-EXP-0001" and edited["editedByName"] == "Admin One"
    cancelled = exp.CancelExpense(db, ExpenseId=made["expenseId"], Reason="duplicate", Actor=admin)
    assert cancelled["status"] == "CANCELLED"
    with pytest.raises(HTTPException) as excinfo:
        exp.CancelExpense(db, ExpenseId=made["expenseId"], Reason="again", Actor=admin)
    assert _code(excinfo) == "EXPENSE_ALREADY_CANCELLED"
    actions = [row.action for row in db.query(PaymentAuditLog).filter(PaymentAuditLog.entity_type == "EXPENSE").order_by(PaymentAuditLog.created_at)]
    assert actions == ["CREATE", "UPDATE", "CANCEL"]


def test_expense_list_totals_filters_and_excel():
    db, admin, monthly, bag = _setup()
    rent, printing = _category(db, "Rent"), _category(db, "Printing")
    for key, day, cat, item, method, amount in [
        ("exp-00000001", "2026-11-01", rent, "Rent", "UPI", 1200000),
        ("exp-00000002", "2026-11-10", printing, "Sheets", "CASH", 50000),
        ("exp-00000003", "2026-10-28", printing, "Posters", "CASH", 30000),
        ("exp-00000004", "2026-11-12", printing, "Mistake", "CASH", 99900),
    ]:
        made = exp.CreateExpense(db, Request={"expenseDate": day, "categoryId": cat.id, "item": item, "methods": [{"method": method, "amountPaise": amount}]}, IdempotencyKey=key, Actor=admin)
    exp.CancelExpense(db, ExpenseId=made["expenseId"], Reason="wrong", Actor=admin)
    november = exp.ListExpenses(db, Filters={"month": "2026-11"})
    assert november["totalCount"] == 3 and november["totals"]["totalDisplay"] == "₹12,500.00"
    assert {row["categoryName"]: row["amountDisplay"] for row in november["totals"]["byCategory"]} == {"Rent": "₹12,000.00", "Printing": "₹500.00"}
    assert {row["methodLabel"]: row["amountDisplay"] for row in november["totals"]["byMethod"]} == {"UPI": "₹12,000.00", "Cash": "₹500.00"}
    assert exp.ListExpenses(db, Filters={"categoryId": printing.id})["totals"]["totalDisplay"] == "₹800.00"
    assert exp.ListExpenses(db, Filters={"search": "poster"})["totalCount"] == 1
    assert exp.ListExpenses(db, Filters={"status": "CANCELLED"})["totalCount"] == 1
    book = load_workbook(io.BytesIO(BuildExpensesWorkbook(exp.ExpensesForExport(db, Filters={"month": "2026-11"}))))
    assert book["Expenses"].cell(row=book["Expenses"].max_row, column=10).value == 12500.0
    assert [tuple(r) for r in book["By Category"].iter_rows(values_only=True)][-1] == ("Total", 12500.0)


# --- API -----------------------------------------------------------------------

def test_reports_and_expenses_api():
    from app.main import app

    db, admin, monthly, bag = _world_with_payments()
    app.dependency_overrides[get_current_user] = lambda: admin
    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)
    try:
        assert client.get("/api/admin/payments/reports/overview").json()["due"]["display"] == "₹2,300.00"
        coll = client.get("/api/admin/payments/reports/collections", params={"dateFrom": "2026-11-01", "dateTo": "2026-11-30"}).json()
        assert coll["total"]["display"] == "₹1,500.00"
        xlsx = client.get("/api/admin/payments/reports/collections/export", params={"dateFrom": "2026-11-20"})
        assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK" and "2026-11-20" in xlsx.headers["content-disposition"]
        pdf = client.get("/api/admin/payments/reports/collections/pdf", params={"dateFrom": "2026-11-20", "method": "CASH", "receivedBy": "admin-1"})
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
        assert "Cash only" in PdfReader(io.BytesIO(pdf.content)).pages[0].extract_text()
        dues = client.get("/api/admin/payments/reports/dues", params={"bucket": "D0_30"}).json()
        assert dues["total"]["display"] == "₹1,400.00"
        assert client.get("/api/admin/payments/reports/dues/export").content[:2] == b"PK"
        cats = client.get("/api/admin/payments/expense-categories").json()["categories"]
        rent = next(row for row in cats if row["name"] == "Rent")
        made = client.post("/api/admin/payments/expenses", json={"idempotencyKey": "api-exp-0001", "categoryId": rent["categoryId"], "item": "Rent", "expenseDate": "2026-11-02", "methods": [{"method": "UPI", "amountPaise": 100000}]})
        assert made.status_code == 200 and made.json()["expenseNumber"] == "MP-EXP-0001"
        expense_id = made.json()["expenseId"]
        assert client.get("/api/admin/payments/expenses", params={"month": "2026-11"}).json()["totals"]["totalDisplay"] == "₹1,000.00"
        assert client.get(f"/api/admin/payments/expenses/{expense_id}").json()["item"] == "Rent"
        edit = client.put(f"/api/admin/payments/expenses/{expense_id}", json={"reason": "fix", "categoryId": rent["categoryId"], "item": "Rent Nov", "methods": [{"method": "UPI", "amountPaise": 100000}]})
        assert edit.status_code == 200 and edit.json()["item"] == "Rent Nov"
        assert client.get("/api/admin/payments/expenses/export").content[:2] == b"PK"
        assert client.post(f"/api/admin/payments/expenses/{expense_id}/cancel", json={"reason": "x"}).json()["status"] == "CANCELLED"
        new_cat = client.post("/api/admin/payments/expense-categories", json={"name": "Exam Fees"})
        assert new_cat.status_code == 200
        assert client.patch(f"/api/admin/payments/expense-categories/{new_cat.json()['categoryId']}", json={"name": "Exam Fee"}).json()["name"] == "Exam Fee"
        assert client.get("/api/admin/payments/expenses/nope").status_code == 404
    finally:
        app.dependency_overrides.clear()
