"""Birthdays (2026-10-18): reading dates of birth, picking wishes, the
morning job and the student's greeting."""

from datetime import date, datetime, timezone

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base
from app.models import BirthdayGreeting, BirthdayWish, Level, Module, Notification, Student, Teacher, User
from app.services import birthday_service as B
from app.services.birthday_wishes_seed import SEED_WISHES, SeedBirthdayWishes


@pytest.fixture()
def db():
    engine = create_engine("sqlite://", future=True, connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)()
    session.add(Module(id="m1", module_code="IM", module_name="Intermediate Module"))
    session.add(Level(id="l1", module_id="m1", level_code="IM-L2", level_name="IM Level 2"))
    session.add(User(id="tu", full_name="Meera Iyer", email="t@x.test", password_hash="x", role="TEACHER", is_active=True))
    session.add(Teacher(id="t1", user_id="tu", teacher_code="MP-TE-001", is_active=True))
    session.commit()
    SeedBirthdayWishes(session)
    yield session
    session.close()


def _student(db, n, dob, *, active=True, demo=False, level="l1", teacher="t1", name=None):
    db.add(User(id=f"u{n}", full_name=name or f"Student {n:03d} Test", email=f"s{n}@x.test", password_hash="x", role="STUDENT", is_active=active))
    db.add(Student(id=f"s{n}", user_id=f"u{n}", student_code=f"MP-ST-{n:03d}", dob=dob, is_active=active, is_demo=demo,
                   current_level_id=level, teacher_id=teacher))
    db.commit()
    return db.query(Student).filter(Student.id == f"s{n}").one()


IST_8AM_18_OCT = datetime(2026, 10, 18, 2, 30, tzinfo=timezone.utc)
IST_6AM_18_OCT = datetime(2026, 10, 18, 0, 30, tzinfo=timezone.utc)


# ---------- dates ------------------------------------------------------------

@pytest.mark.parametrize("text,expected", [
    ("2016-11-16", date(2016, 11, 16)),
    ("16-11-2016", date(2016, 11, 16)),
    ("16/11/2016", date(2016, 11, 16)),
    ("2016/11/16", date(2016, 11, 16)),
    ("2016-11-16 00:00:00", date(2016, 11, 16)),
    ("16.11.2016", date(2016, 11, 16)),
    ("", None), (None, None), ("16/11", None), ("not a date", None), ("2016-02-30", None),
])
def test_parse_dob(text, expected):
    assert B.ParseDob(text) == expected


def test_dob_status():
    today = date(2026, 10, 18)
    assert B.DobStatus("2016-11-16", today) == "OK"
    assert B.DobStatus("", today) == "MISSING"
    assert B.DobStatus("sometime", today) == "UNREADABLE"
    assert B.DobStatus("2027-01-01", today) == "UNLIKELY"   # in the future
    assert B.DobStatus("2025-01-01", today) == "UNLIKELY"   # 1 year old
    assert B.DobStatus("1990-01-01", today) == "UNLIKELY"   # 36 years old


def test_29_february_is_celebrated_on_28_february_in_other_years():
    dob = date(2016, 2, 29)
    assert B.BirthdayOn(dob, 2027) == date(2027, 2, 28)
    assert B.BirthdayOn(dob, 2028) == date(2028, 2, 29)
    assert B.BirthdayOn(date(2016, 3, 1), 2027) == date(2027, 3, 1)


def test_age_bands():
    assert [B.AgeBand(a) for a in (5, 8, 9, 12, 13, 17)] == ["YOUNG", "YOUNG", "MIDDLE", "MIDDLE", "OLDER", "OLDER"]


def test_only_active_students_with_a_birthday_that_day(db):
    _student(db, 1, "2016-10-18")
    _student(db, 2, "18-10-2015")
    _student(db, 3, "2016-10-18", active=False)
    _student(db, 4, "2016-10-18", demo=True)
    _student(db, 5, "2016-10-19")
    _student(db, 6, "")
    rows = B.BirthdayStudentsOn(db, date(2026, 10, 18))
    assert [r[0].id for r in rows] == ["s1", "s2"]
    assert [r[3] for r in rows] == [10, 11]
    with_demo = B.BirthdayStudentsOn(db, date(2026, 10, 18), IncludeDemo=True)
    assert [r[0].id for r in with_demo] == ["s1", "s2", "s4"]


# ---------- wishes -----------------------------------------------------------

def test_the_seed_has_about_sixty_unique_wishes_and_never_duplicates(db):
    assert len(SEED_WISHES) >= 60
    assert len({text for *_, text in SEED_WISHES}) == len(SEED_WISHES)
    assert db.query(BirthdayWish).count() == len(SEED_WISHES)
    wish = db.query(BirthdayWish).filter(BirthdayWish.seed_key == "warm-01").one()
    wish.text = "Edited by an admin, {name}."
    wish.is_active = False
    db.commit()
    assert SeedBirthdayWishes(db) == 0
    wish = db.query(BirthdayWish).filter(BirthdayWish.seed_key == "warm-01").one()
    assert wish.text == "Edited by an admin, {name}." and wish.is_active is False


def test_every_seed_wish_renders_cleanly():
    for _key, style, band, text in SEED_WISHES:
        assert style in B.STYLES and band in B.AGE_BANDS
        out = B.RenderWish(text, Name="Aarav", Age=9, Level="IM-L2")
        assert "{" not in out and "}" not in out
        assert "  " not in out


def test_no_two_students_get_the_same_wish_on_the_same_day(db):
    for n in range(1, 13):
        _student(db, n, f"2016-10-18")
    B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT)
    wishes = [g.wish_id for g in db.query(BirthdayGreeting).all()]
    assert len(wishes) == 12 and len(set(wishes)) == 12


def test_no_repeat_for_a_student_within_three_years(db):
    s = _student(db, 1, "2016-10-18")
    got = []
    for year in range(2026, 2030):
        when = datetime(year, 10, 18, 2, 30, tzinfo=timezone.utc)
        B.EnsureBirthdayGreetings(db, when)
        got.append(db.query(BirthdayGreeting).filter(BirthdayGreeting.year == year).one().wish_id)
    assert len(set(got)) == 4


def test_age_band_and_level_rules(db):
    young = _student(db, 1, "2020-10-18")           # turns 6
    no_level = _student(db, 2, "2012-10-18", level=None)  # turns 14
    B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT)
    g1 = db.query(BirthdayGreeting).filter(BirthdayGreeting.student_id == "s1").one()
    g2 = db.query(BirthdayGreeting).filter(BirthdayGreeting.student_id == "s2").one()
    w1 = db.get(BirthdayWish, g1.wish_id)
    w2 = db.get(BirthdayWish, g2.wish_id)
    assert w1.age_band in ("ALL", "YOUNG")
    assert w2.age_band in ("ALL", "OLDER") and not w2.needs_level
    assert "Student" in g1.rendered_text or "{" not in g1.rendered_text
    assert g1.age == 6 and g2.age == 14


def test_falls_back_when_every_fitting_wish_is_taken(db):
    db.query(BirthdayWish).update({BirthdayWish.is_active: False})
    db.query(BirthdayWish).filter(BirthdayWish.seed_key == "short-01").update({BirthdayWish.is_active: True})
    db.commit()
    _student(db, 1, "2016-10-18")
    _student(db, 2, "2016-10-18")
    B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT)
    rows = db.query(BirthdayGreeting).all()
    assert len(rows) == 2 and all(r.wish_id for r in rows)


def test_still_kind_when_every_wish_is_switched_off(db):
    db.query(BirthdayWish).update({BirthdayWish.is_active: False})
    db.commit()
    _student(db, 1, "2016-10-18", name="Aarav Shah")
    B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT)
    g = db.query(BirthdayGreeting).one()
    assert g.wish_id is None and "Aarav" in g.rendered_text


# ---------- the morning job --------------------------------------------------

def test_nothing_before_seven_india_time(db):
    _student(db, 1, "2016-10-18")
    assert B.EnsureBirthdayGreetings(db, IST_6AM_18_OCT) == 0
    assert db.query(Notification).count() == 0


def test_job_is_safe_to_run_again(db):
    _student(db, 1, "2016-10-18", name="Aarav Shah")
    _student(db, 2, "2016-10-18")
    assert B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT) == 2
    assert B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT) == 0
    assert db.query(BirthdayGreeting).count() == 2
    notes = db.query(Notification).filter(Notification.type == "BIRTHDAY_GREETING").all()
    assert len(notes) == 2
    first = [n for n in notes if n.recipient_user_id == "u1"][0]
    assert first.title == "Happy Birthday, Aarav!"
    assert first.category == "BIRTHDAY" and first.color_variant == "ROSE"
    assert first.target_route == "/student/dashboard?birthday=1"


def test_greeting_made_early_on_the_dashboard_is_kept_and_notified_once(db):
    s = _student(db, 1, "2016-10-18")
    early = B.GreetingFor(db, s, date(2026, 10, 18))
    assert early is not None and early.notified_student_at is None
    assert B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT) == 1
    again = db.query(BirthdayGreeting).one()
    assert again.id == early.id and again.rendered_text == early.rendered_text
    assert again.notified_student_at is not None


def test_leap_day_birthday_on_28_february(db):
    s = _student(db, 1, "2016-02-29")
    when = datetime(2027, 2, 28, 3, 0, tzinfo=timezone.utc)
    assert B.EnsureBirthdayGreetings(db, when) == 1
    g = db.query(BirthdayGreeting).one()
    assert g.birthday_date == date(2027, 2, 28) and g.age == 11
    assert B.GreetingFor(db, s, date(2027, 3, 1)) is None


def test_the_day_follows_india_time(db):
    _student(db, 1, "2016-10-19")
    late_utc = datetime(2026, 10, 18, 20, 0, tzinfo=timezone.utc)   # 19 Oct, 01:30 IST
    assert B.EnsureBirthdayGreetings(db, late_utc) == 0             # before 07:00 IST
    morning = datetime(2026, 10, 19, 1, 45, tzinfo=timezone.utc)    # 19 Oct, 07:15 IST
    assert B.EnsureBirthdayGreetings(db, morning) == 1


def test_new_year_birthdays(db):
    _student(db, 1, "2016-01-01")
    _student(db, 2, "2015-12-31")
    assert B.EnsureBirthdayGreetings(db, datetime(2026, 12, 31, 2, 0, tzinfo=timezone.utc)) == 1
    assert B.EnsureBirthdayGreetings(db, datetime(2027, 1, 1, 2, 0, tzinfo=timezone.utc)) == 1
    years = sorted((g.student_id, g.year) for g in db.query(BirthdayGreeting).all())
    assert years == [("s1", 2027), ("s2", 2026)]


# ---------- the student's greeting -------------------------------------------

def test_greeting_payload_and_marks(db):
    s = _student(db, 1, "2016-10-18", name="Aarav Shah")
    g = B.GreetingFor(db, s, date(2026, 10, 18))
    payload = B.GreetingPayload(db, s, g)
    assert payload["firstName"] == "Aarav" and payload["age"] == 10
    assert payload["signedBy"] == "Meera Iyer and everyone at MathPath"
    assert payload["seen"] is False and payload["dismissed"] is False
    B.MarkGreeting(db, s, g.id, Seen=True)
    B.MarkGreeting(db, s, g.id, Dismissed=True)
    payload = B.GreetingPayload(db, s, db.get(BirthdayGreeting, g.id))
    assert payload["seen"] is True and payload["dismissed"] is True


def test_a_student_cannot_mark_someone_elses_greeting(db):
    s1 = _student(db, 1, "2016-10-18")
    s2 = _student(db, 2, "2016-10-18")
    g = B.GreetingFor(db, s1, date(2026, 10, 18))
    assert B.MarkGreeting(db, s2, g.id, Dismissed=True) is None
    assert db.get(BirthdayGreeting, g.id).dismissed_at is None


def test_no_greeting_on_other_days_or_with_a_bad_dob(db):
    s = _student(db, 1, "2016-10-19")
    bad = _student(db, 2, "18/10")
    assert B.GreetingFor(db, s, date(2026, 10, 18)) is None
    assert B.GreetingFor(db, bad, date(2026, 10, 18)) is None


def test_signed_by_everyone_when_no_teacher(db):
    s = _student(db, 1, "2016-10-18", teacher=None)
    assert B.SignedBy(db, s) == "Everyone at MathPath"


# ---------- staff views (Part 2) ---------------------------------------------

NOON_18_OCT = datetime(2026, 10, 18, 6, 30, tzinfo=timezone.utc)   # 12:00 IST


def test_staff_list_range_order_counts_and_status(db):
    _student(db, 1, "2016-10-18", name="Zara Khan")
    _student(db, 2, "2015-10-18", name="Aarav Shah")
    _student(db, 3, "2014-10-20")
    _student(db, 4, "2013-11-30")
    _student(db, 5, "2016-10-05")            # earlier this month
    _student(db, 6, "")                      # missing
    _student(db, 7, "2016-10-18", demo=True)  # hidden by default
    B.EnsureBirthdayGreetings(db, NOON_18_OCT)
    data = B.StaffBirthdays(db, Start=date(2026, 10, 18), End=date(2026, 11, 16), Now=NOON_18_OCT)
    names = [(e["date"], e["name"]) for e in data["entries"]]
    assert names == [("2026-10-18", "Aarav Shah"), ("2026-10-18", "Zara Khan"), ("2026-10-20", "Student 003 Test")]
    assert data["counts"] == {"today": 2, "thisWeek": 3, "thisMonth": 4, "missingDob": 1}
    today = data["entries"][0]
    assert today["isToday"] and today["status"] == "SENT" and today["age"] == 11 and today["wish"]
    assert today["levelCode"] == "IM-L2" and today["teacherName"] == "Meera Iyer"
    assert data["entries"][2]["status"] == "SCHEDULED"


def test_staff_list_crosses_the_new_year_and_marks_missed_days(db):
    _student(db, 1, "2016-12-30")
    _student(db, 2, "2016-01-02")
    _student(db, 3, "2016-10-10")
    data = B.StaffBirthdays(db, Start=date(2026, 12, 25), End=date(2027, 1, 5), Now=NOON_18_OCT)
    assert [e["date"] for e in data["entries"]] == ["2026-12-30", "2027-01-02"]
    past = B.StaffBirthdays(db, Start=date(2026, 10, 1), End=date(2026, 10, 31), Now=NOON_18_OCT)
    assert [e["status"] for e in past["entries"]] == ["NOT_SENT"]


def test_staff_filters_and_teacher_scope(db):
    db.add(User(id="tu2", full_name="Ravi Das", email="t2@x.test", password_hash="x", role="TEACHER", is_active=True))
    db.add(Teacher(id="t2", user_id="tu2", teacher_code="MP-TE-002", is_active=True))
    db.commit()
    _student(db, 1, "2016-10-20", name="Aarav Shah")
    _student(db, 2, "2016-10-21", teacher="t2", level=None)
    _student(db, 3, "2016-10-22", teacher=None)
    span = dict(Start=date(2026, 10, 18), End=date(2026, 10, 31), Now=NOON_18_OCT)
    ids = lambda **kw: [e["studentId"] for e in B.StaffBirthdays(db, **span, **kw)["entries"]]
    assert ids() == ["s1", "s2", "s3"]
    assert ids(TeacherId="t2") == ["s2"]
    assert ids(TeacherId="NONE") == ["s3"]
    assert ids(LevelId="l1") == ["s1", "s3"]
    assert ids(Search="aarav") == ["s1"]
    assert ids(Search="MP-ST-002") == ["s2"]
    # a teacher's own view ignores any other teacher filter
    assert ids(OnlyTeacherId="t2", TeacherId="t1") == ["s2"]


def test_show_demo_switch_includes_demo_students(db):
    from app.services.demo_students import SetIncludeDemo
    _student(db, 1, "2016-10-20", demo=True)
    span = dict(Start=date(2026, 10, 18), End=date(2026, 10, 31), Now=NOON_18_OCT)
    assert B.StaffBirthdays(db, **span)["entries"] == []
    SetIncludeDemo(db, True)
    assert [e["studentId"] for e in B.StaffBirthdays(db, **span)["entries"]] == ["s1"]


def test_missing_dob_list(db):
    _student(db, 1, "")
    _student(db, 2, "18/10")
    _student(db, 3, "2030-01-01")
    _student(db, 4, "2016-10-18")
    _student(db, 5, "", active=False)
    rows = B.MissingDob(db)
    assert sorted((r["studentId"], r["problem"]) for r in rows) == [("s1", "MISSING"), ("s2", "UNREADABLE"), ("s3", "UNLIKELY")]


def test_export_has_no_phone_numbers(db):
    from app.services.birthday_export import BuildBirthdaysWorkbook
    from openpyxl import load_workbook
    from io import BytesIO
    _student(db, 1, "2016-10-20", name="Aarav Shah")
    data = B.StaffBirthdays(db, Start=date(2026, 10, 18), End=date(2026, 10, 31), Now=NOON_18_OCT)
    sheet = load_workbook(BytesIO(BuildBirthdaysWorkbook(data["entries"]))).active
    assert [c.value for c in sheet[1]] == ["Date", "Day", "Student", "Student ID", "Turns", "Level", "Teacher", "Centre", "Wish"]
    assert sheet.cell(row=2, column=3).value == "Aarav Shah" and sheet.cell(row=2, column=5).value == 10


# ---------- Part 3: admin summary, managing wishes ------------------------------

def _admin(db, n=1, active=True):
    db.add(User(id=f"a{n}", full_name=f"Admin {n}", email=f"a{n}@x.test", password_hash="x", role="ADMIN", is_active=active))
    db.commit()


def test_admins_get_one_summary_a_day_without_demo_students(db):
    _admin(db, 1)
    _admin(db, 2, active=False)
    _student(db, 1, "2016-10-18", name="Zara Khan")
    _student(db, 2, "2015-10-18", name="Aarav Shah")
    _student(db, 3, "2016-10-18", demo=True, name="Demo Kid")
    B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT)
    B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT)
    rows = db.query(Notification).filter(Notification.type == "BIRTHDAYS_TODAY").all()
    assert len(rows) == 1 and rows[0].recipient_user_id == "a1"
    assert rows[0].id == B.AdminSummaryId("a1", date(2026, 10, 18))
    assert rows[0].title == "2 Birthdays Today"
    assert rows[0].message.startswith("Aarav Shah (turns 11), Zara Khan (turns 10)")
    assert "Demo" not in rows[0].message and rows[0].target_route == "/admin/birthdays"


def test_admin_summary_updates_when_a_birthday_is_added_later(db):
    _admin(db, 1)
    _student(db, 1, "2016-10-18", name="Zara Khan")
    B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT)
    late = _student(db, 2, "", name="Aarav Shah")
    late.dob = "2015-10-18"
    db.commit()
    B.EnsureBirthdayGreetings(db, NOON_18_OCT)
    rows = db.query(Notification).filter(Notification.type == "BIRTHDAYS_TODAY").all()
    assert len(rows) == 1 and rows[0].title == "2 Birthdays Today"


def test_no_admin_summary_for_only_demo_birthdays(db):
    _admin(db, 1)
    _student(db, 1, "2016-10-18", demo=True)
    B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT)
    assert db.query(Notification).filter(Notification.type == "BIRTHDAYS_TODAY").count() == 0


def test_list_wishes_with_usage(db):
    for n in range(1, 4):
        _student(db, n, "2016-10-18")
    B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT)
    data = B.ListWishes(db)
    assert len(data["wishes"]) == len(SEED_WISHES)
    used = [w for w in data["wishes"] if w["uses"]]
    assert len(used) == 3 and all(w["lastUsed"] == "2026-10-18" for w in used)
    assert data["activeByAge"]["YOUNG"] >= 40 and data["minimumPerAge"] == 10
    assert [w["style"] for w in data["wishes"]][0] == "WARM"


def test_create_wish_validates_and_is_then_used(db):
    with pytest.raises(B.WishError):
        B.CreateWish(db, Text="  ", Style="WARM", AgeBand="ALL", UserId="a1")
    with pytest.raises(B.WishError, match="isn't something"):
        B.CreateWish(db, Text="Hi {nam}!", Style="WARM", AgeBand="ALL", UserId="a1")
    with pytest.raises(B.WishError, match="already in the set"):
        B.CreateWish(db, Text=SEED_WISHES[3][3].upper(), Style="WARM", AgeBand="ALL", UserId="a1")
    with pytest.raises(B.WishError, match="Choose a style"):
        B.CreateWish(db, Text="New one, {name}.", Style="LOUD", AgeBand="ALL", UserId="a1")
    wish = B.CreateWish(db, Text="  Have a  lovely day in {level}, {name}! ", Style="PROUD", AgeBand="OLDER", UserId="a1")
    assert wish.text == "Have a lovely day in {level}, {name}!" and wish.needs_level is True and wish.seed_key is None


def test_editing_a_wish_never_changes_what_a_student_read(db):
    _student(db, 1, "2016-10-18", name="Aarav Shah")
    B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT)
    g = db.query(BirthdayGreeting).one()
    before = g.rendered_text
    B.UpdateWish(db, g.wish_id, Text="Totally new words, {name}.", UserId="a1")
    db.expire_all()
    assert db.get(BirthdayGreeting, g.id).rendered_text == before
    assert db.get(BirthdayWish, g.wish_id).text == "Totally new words, {name}."


def test_switching_off_keeps_ten_per_age_group(db):
    # leave exactly 10 wishes a 13+ student could get, then try to switch one off
    keep = [w for w in db.query(BirthdayWish).all() if w.age_band in ("ALL", "OLDER")][:10]
    keep_ids = {w.id for w in keep}
    for w in db.query(BirthdayWish).all():
        if w.age_band in ("ALL", "OLDER") and w.id not in keep_ids:
            w.is_active = False
    db.commit()
    assert B.ActiveWishCounts(db)["OLDER"] == 10
    with pytest.raises(B.WishError, match="fewer than 10 wishes for ages 13 and over"):
        B.UpdateWish(db, keep[0].id, IsActive=False, UserId="a1")
    assert db.get(BirthdayWish, keep[0].id).is_active is True
    young_only = db.query(BirthdayWish).filter(BirthdayWish.age_band == "YOUNG").first()
    B.UpdateWish(db, young_only.id, IsActive=False, UserId="a1")  # a 3-8 wish: fine
    assert db.get(BirthdayWish, young_only.id).is_active is False


def test_preview_fills_in_samples():
    out = B.PreviewWish("Happy {age}th, {name}! Enjoy {level}.")
    assert [s["text"] for s in out][1] == "Happy 11th, Diya! Enjoy IM-L2."


# ---------- Part 4: teachers ------------------------------------------------------

def test_teacher_hears_about_their_own_students_once(db):
    db.add(User(id="tu2", full_name="Ravi Das", email="t2@x.test", password_hash="x", role="TEACHER", is_active=True))
    db.add(Teacher(id="t2", user_id="tu2", teacher_code="MP-TE-002", is_active=True))
    db.commit()
    _student(db, 1, "2016-10-18", name="Aarav Shah")                # Meera's
    _student(db, 2, "2015-10-18", name="Diya Nair", teacher="t2")   # Ravi's
    _student(db, 3, "2016-10-18", name="Demo Kid", demo=True)       # demo: no teacher notice
    _student(db, 4, "2016-10-18", name="No Teacher", teacher=None)
    B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT)
    B.EnsureBirthdayGreetings(db, NOON_18_OCT)
    rows = db.query(Notification).filter(Notification.type == "STUDENT_BIRTHDAY").all()
    got = sorted((r.recipient_user_id, r.title) for r in rows)
    assert got == [("tu", "Aarav Shah Turns 10 Today"), ("tu2", "Diya Nair Turns 11 Today")]
    assert all(r.target_route == "/teacher/birthdays" and r.category == "BIRTHDAY" for r in rows)
    greetings = {g.student_id: g for g in db.query(BirthdayGreeting).all()}
    assert greetings["s1"].notified_teacher_at and greetings["s4"].notified_teacher_at
    assert greetings["s3"].notified_teacher_at is None


def test_no_teacher_notice_for_an_inactive_teacher(db):
    db.query(Teacher).filter(Teacher.id == "t1").update({Teacher.is_active: False})
    db.commit()
    _student(db, 1, "2016-10-18")
    B.EnsureBirthdayGreetings(db, IST_8AM_18_OCT)
    assert db.query(Notification).filter(Notification.type == "STUDENT_BIRTHDAY").count() == 0


def test_teacher_level_options(db):
    db.add(Level(id="l2", module_id="m1", level_code="IM-L3", level_name="IM Level 3"))
    db.commit()
    _student(db, 1, "2016-10-18")
    _student(db, 2, "2016-10-18", level="l2", teacher=None)
    assert [l["code"] for l in B.TeacherLevelOptions(db, "t1")] == ["IM-L2"]


# ---------- Part 5: the clock ---------------------------------------------------

def _ist(y, m, d, hh, mm, ss=0):
    """A moment given in India time, as the UTC time the server would see."""
    return datetime(y, m, d, hh, mm, ss, tzinfo=B.INDIA).astimezone(timezone.utc)


def test_the_send_time_is_exactly_seven(db):
    _student(db, 1, "2016-10-18")
    assert B.EnsureBirthdayGreetings(db, _ist(2026, 10, 18, 6, 59, 59)) == 0
    early = B.StaffBirthdays(db, Start=date(2026, 10, 18), End=date(2026, 10, 18), Now=_ist(2026, 10, 18, 6, 59, 59))
    assert early["entries"][0]["status"] == "SCHEDULED"
    assert B.EnsureBirthdayGreetings(db, _ist(2026, 10, 18, 7, 0, 0)) == 1
    after = B.StaffBirthdays(db, Start=date(2026, 10, 18), End=date(2026, 10, 18), Now=_ist(2026, 10, 18, 7, 0, 30))
    assert after["entries"][0]["status"] == "SENT"
    # the rest of the day sends nothing more, right up to midnight
    assert B.EnsureBirthdayGreetings(db, _ist(2026, 10, 18, 23, 59, 59)) == 0
    assert db.query(Notification).filter(Notification.type == "BIRTHDAY_GREETING").count() == 1


def test_midnight_india_time_moves_today_while_utc_is_still_yesterday(db):
    _student(db, 1, "2016-10-18", name="Zara Khan")
    _student(db, 2, "2016-10-19", name="Aarav Shah")
    just_after = _ist(2026, 10, 19, 0, 5)
    assert just_after.date() == date(2026, 10, 18)          # UTC still says the 18th
    data = B.StaffBirthdays(db, Start=date(2026, 10, 19), End=date(2026, 10, 25), Now=just_after)
    assert data["counts"]["today"] == 1
    assert [(e["name"], e["isToday"], e["status"]) for e in data["entries"]] == [("Aarav Shah", True, "SCHEDULED")]
    s2 = db.query(Student).filter(Student.id == "s2").one()
    assert B.IsBirthdayToday(s2, date(2026, 10, 19)) and not B.IsBirthdayToday(s2, date(2026, 10, 18))


def test_a_missed_morning_is_not_sent_late_the_next_day(db):
    _student(db, 1, "2016-10-18")
    assert B.EnsureBirthdayGreetings(db, _ist(2026, 10, 19, 9, 0)) == 0      # server was down on the 18th
    assert db.query(BirthdayGreeting).count() == 0
    data = B.StaffBirthdays(db, Start=date(2026, 10, 1), End=date(2026, 10, 31), Now=_ist(2026, 10, 19, 9, 0))
    assert data["entries"][0]["status"] == "NOT_SENT"


def test_leap_day_in_a_leap_year_is_on_29_february(db):
    _student(db, 1, "2016-02-29")
    assert B.EnsureBirthdayGreetings(db, _ist(2028, 2, 28, 8, 0)) == 0
    assert B.EnsureBirthdayGreetings(db, _ist(2028, 2, 29, 8, 0)) == 1
    g = db.query(BirthdayGreeting).one()
    assert g.birthday_date == date(2028, 2, 29) and g.age == 12
    feb = B.StaffBirthdays(db, Start=date(2027, 2, 1), End=date(2027, 2, 28), Now=_ist(2027, 1, 10, 12, 0))
    assert [e["date"] for e in feb["entries"]] == ["2027-02-28"]


def test_new_year_eve_and_day_by_india_time(db):
    _student(db, 1, "2016-01-01", name="Ira Rao")
    _student(db, 2, "2015-12-31", name="Kabir Das")
    eve = _ist(2026, 12, 31, 23, 59)
    day = _ist(2027, 1, 1, 0, 1)
    assert [e["name"] for e in B.StaffBirthdays(db, Start=date(2026, 12, 31), End=date(2026, 12, 31), Now=eve)["entries"]] == ["Kabir Das"]
    week = B.StaffBirthdays(db, Start=date(2027, 1, 1), End=date(2027, 1, 7), Now=day)
    assert week["counts"]["today"] == 1 and week["entries"][0]["name"] == "Ira Rao" and week["entries"][0]["age"] == 11
    assert B.EnsureBirthdayGreetings(db, day) == 0                 # still before 07:00
    assert B.EnsureBirthdayGreetings(db, _ist(2027, 1, 1, 7, 0)) == 1


def test_the_background_check_wakes_on_the_hour():
    assert B.SecondsUntilNextCheck(_ist(2026, 10, 18, 6, 50)) == 10 * 60 + 5
    assert B.SecondsUntilNextCheck(_ist(2026, 10, 18, 6, 59, 59)) == 6
    assert B.SecondsUntilNextCheck(_ist(2026, 10, 18, 7, 0, 10)) == 30 * 60
    assert B.SecondsUntilNextCheck(_ist(2026, 10, 18, 3, 0)) == 30 * 60
    assert B.SecondsUntilNextCheck(_ist(2026, 10, 18, 23, 59)) == 30 * 60
