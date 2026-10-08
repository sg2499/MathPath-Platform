"""Payments Phase 3 (2026-10-08): payments at the counter -- part-payments,
split methods, discount, advances, edit, cancel, receipts, the student's
account, PDF, Excel and the API."""

import io
import json

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.database import get_db
from app.dependencies import get_current_user
from app.models import PaymentAllocation, PaymentAuditLog, PaymentInvoice, PaymentNumberSequence, PaymentReceipt, User
from app.services.payments import invoice_pdf, numbering
from app.services.payments import invoices_service as inv
from app.services.payments import receipts_service as rec
from app.services.payments.invoice_export import BuildPaymentsWorkbook
from tests.test_payments_phase2_invoices import _code, _req, _session, _world

KEY = iter(f"pay-key-{n:06d}" for n in range(100000))


def _setup(receipts=True):
    db = _session()
    admin, monthly, bag = _world(db)
    if receipts:
        numbering.SetStartingNumber(db, Key="RECEIPT", NextNumber=631, Actor=admin)
        db.commit()
    inv.GenerateInvoices(db, Request=_req([monthly, bag], students=("s1", "s2")), IdempotencyKey="gen-00000001", Actor=admin)
    return db, admin, monthly, bag


def _invoice(db, student, fee):
    return db.query(PaymentInvoice).filter(PaymentInvoice.student_id == student, PaymentInvoice.fee_name == fee).one()


def _pay(db, admin, student="s1", allocations=(), methods=(("CASH", 110000, None),), key=None, **extra):
    request = {
        "paymentDate": "2026-10-08",
        "payBy": "Parent",
        "allocations": [{"invoiceId": i, "amountPaise": a, "discountPaise": d} for i, a, d in allocations],
        "methods": [{"method": m, "amountPaise": a, "reference": r} for m, a, r in methods],
        **extra,
    }
    return rec.RecordPayment(db, StudentId=student, Request=request, IdempotencyKey=key or next(KEY), Actor=admin)


@pytest.fixture(autouse=True)
def _fixed_today(monkeypatch):
    from datetime import date

    monkeypatch.setattr(inv, "TodayInIndia", lambda: date(2026, 11, 20))


# --- record --------------------------------------------------------------------

def test_full_payment_settles_the_invoice_and_issues_a_receipt():
    db, admin, monthly, bag = _setup()
    fee = _invoice(db, "s1", "Monthly Fee")
    result = _pay(db, admin, allocations=[(fee.id, 110000, 0)])
    assert result["receiptNumber"] == "MP-MRCPT-631" and result["amountDisplay"] == "₹1,100.00"
    assert result["receivedByName"] == "Admin One" and result["methodSummary"] == "Cash ₹1,100.00"
    db.refresh(fee)
    assert (fee.status, fee.paid_paise, fee.discount_paise) == ("PAID", 110000, 0)
    assert db.get(PaymentNumberSequence, "RECEIPT").next_number == 632
    entry = db.query(PaymentAuditLog).filter(PaymentAuditLog.entity_type == "PAYMENT", PaymentAuditLog.action == "CREATE").one()
    assert json.loads(entry.after_json)["appliedTo"].startswith("MP-INV-0010")


def test_part_payments_and_split_methods():
    db, admin, monthly, bag = _setup()
    fee = _invoice(db, "s1", "Monthly Fee")
    _pay(db, admin, allocations=[(fee.id, 50000, 0)], methods=[("CASH", 20000, None), ("UPI", 30000, "UTR4312")])
    db.refresh(fee)
    assert fee.status == "PART_PAID" and inv.InvoicePayload(fee)["balanceDisplay"] == "₹600.00"
    second = _pay(db, admin, allocations=[(fee.id, 60000, 0)], methods=[("UPI", 60000, "UTR9")])
    db.refresh(fee)
    assert fee.status == "PAID" and second["receiptNumber"] == "MP-MRCPT-632"


def test_one_payment_can_cover_several_invoices():
    db, admin, monthly, bag = _setup()
    fee, bagi = _invoice(db, "s1", "Monthly Fee"), _invoice(db, "s1", "MathPath Bag")
    _pay(db, admin, allocations=[(fee.id, 110000, 0), (bagi.id, 35000, 0)], methods=[("UPI", 145000, "UTR1")])
    db.refresh(fee)
    db.refresh(bagi)
    assert fee.status == bagi.status == "PAID"


@pytest.mark.parametrize(
    "allocations,methods,extra,code",
    [
        ([("FEE", 120000, 0)], [("CASH", 120000, None)], {}, "MORE_THAN_DUE"),
        ([("FEE", 110000, 0)], [("CASH", 100000, None)], {}, "METHODS_SHORT"),
        ([("FEE", 100000, 0)], [("CASH", 110000, None)], {}, "EXTRA_NOT_ALLOWED"),
        ([("FEE", 110000, 0)], [("UPI", 110000, None)], {}, "REFERENCE_REQUIRED"),
        ([("FEE", 110000, 0)], [], {}, "METHOD_REQUIRED"),
        ([("FEE", 110000, 0)], [("BITCOIN", 110000, None)], {}, "METHOD_INVALID"),
        ([("FEE", 110000, 0)], [("RAZORPAY", 110000, None)], {}, "METHOD_INVALID"),
        ([("FEE", 100000, 10000)], [("CASH", 100000, None)], {}, "DISCOUNT_REASON_REQUIRED"),
        ([("FEE", 110000, 0)], [("CASH", 110000, None)], {"paymentDate": "2026-12-01"}, "DATE_IN_FUTURE"),
        ([("FEE", -5, 0)], [("CASH", 110000, None)], {}, "AMOUNT_NEGATIVE"),
        ([("OTHER", 110000, 0)], [("CASH", 110000, None)], {}, "INVOICE_NOT_FOUND"),
        ([("FEE", 110000, 0)], [("CASH", 110000, None)], {"receivedByUserId": "user-s1"}, "RECEIVED_BY_INVALID"),
    ],
)
def test_bad_payments_are_refused_with_a_clear_reason(allocations, methods, extra, code):
    db, admin, monthly, bag = _setup()
    ids = {"FEE": _invoice(db, "s1", "Monthly Fee").id, "OTHER": _invoice(db, "s2", "Monthly Fee").id}
    with pytest.raises(HTTPException) as excinfo:
        _pay(db, admin, allocations=[(ids[i], a, d) for i, a, d in allocations], methods=methods, **extra)
    assert _code(excinfo) == code
    db.rollback()
    assert db.query(PaymentReceipt).count() == 0


def test_receipt_numbering_must_be_set_first():
    db, admin, monthly, bag = _setup(receipts=False)
    fee = _invoice(db, "s1", "Monthly Fee")
    with pytest.raises(HTTPException) as excinfo:
        _pay(db, admin, allocations=[(fee.id, 110000, 0)])
    assert _code(excinfo) == "PAYMENT_NUMBERING_NOT_SET"


def test_a_cancelled_invoice_cannot_be_paid():
    db, admin, monthly, bag = _setup()
    fee = _invoice(db, "s1", "Monthly Fee")
    inv.CancelInvoice(db, InvoiceId=fee.id, Reason="wrong", Actor=admin)
    with pytest.raises(HTTPException) as excinfo:
        _pay(db, admin, allocations=[(fee.id, 110000, 0)])
    assert _code(excinfo) == "INVOICE_CANCELLED"


def test_the_same_save_press_never_records_twice():
    db, admin, monthly, bag = _setup()
    fee = _invoice(db, "s1", "Monthly Fee")
    first = _pay(db, admin, allocations=[(fee.id, 110000, 0)], key="same-press-01")
    again = _pay(db, admin, allocations=[(fee.id, 110000, 0)], key="same-press-01")
    assert again["replayed"] is True and again["paymentId"] == first["paymentId"]
    assert db.query(PaymentReceipt).count() == 1


def test_received_by_can_be_another_admin():
    db, admin, monthly, bag = _setup()
    db.add(User(id="admin-2", full_name="Front Desk", email="fd@example.test", password_hash="x", role="ADMIN", is_active=True))
    db.commit()
    fee = _invoice(db, "s1", "Monthly Fee")
    result = _pay(db, admin, allocations=[(fee.id, 110000, 0)], receivedByUserId="admin-2")
    assert result["receivedByName"] == "Front Desk" and result["createdByName"] == "Admin One"
    assert {row["name"] for row in rec.ListStaff(db)} == {"Admin One", "Front Desk"}


# --- discount --------------------------------------------------------------------

def test_discount_settles_the_invoice_and_shows_everywhere():
    db, admin, monthly, bag = _setup()
    fee = _invoice(db, "s1", "Monthly Fee")
    result = _pay(db, admin, allocations=[(fee.id, 100000, 10000)], methods=[("CASH", 100000, None)], discountReason="sibling concession")
    db.refresh(fee)
    assert (fee.status, fee.paid_paise, fee.discount_paise) == ("PAID", 100000, 10000)
    assert result["discountDisplay"] == "₹100.00" and result["discountReason"] == "sibling concession"
    payload = inv.InvoicePayload(fee)
    assert payload["discountDisplay"] == "₹100.00" and payload["balancePaise"] == 0
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(invoice_pdf.RenderInvoicesPdf([fee]))).pages)
    assert "Discount" in text and "100.00" in text and "PAID" in text


# --- advance ---------------------------------------------------------------------

def test_extra_money_kept_as_advance_is_applied_to_the_next_invoices():
    db, admin, monthly, bag = _setup()
    fee = _invoice(db, "s1", "Monthly Fee")
    paid = _pay(db, admin, allocations=[(fee.id, 110000, 0)], methods=[("UPI", 330000, "UTR-ADV")], keepAdvance=True)
    assert paid["advanceDisplay"] == "₹2,200.00"
    assert rec.StudentAdvanceBalance(db, "s1") == 220000
    # The next two months are invoiced: the preview shows the advance, and
    # the new invoices come out paid.
    preview = inv.PreviewInvoices(db, Request=_req([monthly], students=("s1", "s2"), month=12))
    assert preview["advanceAppliedDisplay"] == "₹1,100.00" and preview["advanceStudents"] == 1
    batch = inv.GenerateInvoices(db, Request=_req([monthly], students=("s1", "s2"), month=12), IdempotencyKey="gen-00000002", Actor=admin)
    assert batch["advanceAppliedDisplay"] == "₹1,100.00"
    december = db.query(PaymentInvoice).filter(PaymentInvoice.student_id == "s1", PaymentInvoice.period_key == "2026-12").one()
    assert december.status == "PAID"
    assert db.query(PaymentInvoice).filter(PaymentInvoice.student_id == "s2", PaymentInvoice.period_key == "2026-12").one().status == "PENDING"
    inv.GenerateInvoices(db, Request=_req([monthly], students=("s1",), month=1, year=2027), IdempotencyKey="gen-00000003", Actor=admin)
    assert rec.StudentAdvanceBalance(db, "s1") == 0
    receipt = rec.GetPayment(db, paid["paymentId"])
    assert [a["kind"] for a in receipt["allocations"]] == ["DIRECT", "ADVANCE", "ADVANCE"]
    assert db.query(PaymentAuditLog).filter(PaymentAuditLog.action == "ADVANCE_APPLIED").count() == 2


def test_advance_only_payment_and_apply_now_to_existing_invoices():
    db, admin, monthly, bag = _setup()
    advance = _pay(db, admin, allocations=[], methods=[("CASH", 50000, None)], keepAdvance=True)
    assert advance["advanceDisplay"] == "₹500.00" and advance["invoiceNumbers"] == []
    applied = rec.ApplyAdvanceNow(db, StudentId="s1", Actor=admin)
    # Oldest invoice first: the monthly fee (generated first) gets the ₹500.
    assert applied["appliedDisplay"] == "₹500.00"
    fee = _invoice(db, "s1", "Monthly Fee")
    assert (fee.status, fee.paid_paise) == ("PART_PAID", 50000)
    with pytest.raises(HTTPException) as excinfo:
        rec.ApplyAdvanceNow(db, StudentId="s1", Actor=admin)
    assert _code(excinfo) == "NOTHING_TO_APPLY"


def test_cancelling_a_paid_invoice_moves_the_money_to_advance():
    db, admin, monthly, bag = _setup()
    bagi = _invoice(db, "s1", "MathPath Bag")
    payment = _pay(db, admin, allocations=[(bagi.id, 35000, 0)], methods=[("CASH", 35000, None)])
    result = inv.CancelInvoice(db, InvoiceId=bagi.id, Reason="returned the bag", Actor=admin)
    assert result["status"] == "CANCELLED" and result["movedToAdvanceDisplay"] == "₹350.00"
    assert rec.StudentAdvanceBalance(db, "s1") == 35000
    assert rec.GetPayment(db, payment["paymentId"])["advanceDisplay"] == "₹350.00"
    entry = db.query(PaymentAuditLog).filter(PaymentAuditLog.entity_type == "INVOICE", PaymentAuditLog.action == "CANCEL").one()
    assert json.loads(entry.after_json)["movedToAdvance"] == "350.00"


# --- edit ------------------------------------------------------------------------

def test_edit_changes_amounts_keeps_the_number_and_records_why():
    db, admin, monthly, bag = _setup()
    fee = _invoice(db, "s1", "Monthly Fee")
    payment = _pay(db, admin, allocations=[(fee.id, 110000, 0)], methods=[("CASH", 110000, None)])
    request = {"paymentDate": "2026-10-07", "payBy": "Mother", "allocations": [{"invoiceId": fee.id, "amountPaise": 60000}], "methods": [{"method": "UPI", "amountPaise": 60000, "reference": "UTR77"}]}
    with pytest.raises(HTTPException) as excinfo:
        rec.EditPayment(db, PaymentId=payment["paymentId"], Request=request, Reason=" ", Actor=admin)
    assert _code(excinfo) == "REASON_REQUIRED"
    edited = rec.EditPayment(db, PaymentId=payment["paymentId"], Request=request, Reason="entered the wrong amount", Actor=admin)
    assert edited["receiptNumber"] == payment["receiptNumber"] and edited["amountDisplay"] == "₹600.00"
    assert edited["methodSummary"] == "UPI ₹600.00" and edited["paymentDate"] == "2026-10-07" and edited["editedByName"] == "Admin One"
    db.refresh(fee)
    assert (fee.status, fee.paid_paise) == ("PART_PAID", 60000)
    entry = db.query(PaymentAuditLog).filter(PaymentAuditLog.action == "UPDATE").one()
    assert json.loads(entry.before_json)["amount"] == "1100.00" and json.loads(entry.after_json)["amount"] == "600.00"
    # The old allocation is kept, released.
    assert db.query(PaymentAllocation).filter(PaymentAllocation.released_at.isnot(None)).count() == 1


def test_edit_can_move_a_payment_to_another_invoice():
    db, admin, monthly, bag = _setup()
    fee, bagi = _invoice(db, "s1", "Monthly Fee"), _invoice(db, "s1", "MathPath Bag")
    payment = _pay(db, admin, allocations=[(fee.id, 35000, 0)], methods=[("CASH", 35000, None)])
    rec.EditPayment(db, PaymentId=payment["paymentId"], Request={"allocations": [{"invoiceId": bagi.id, "amountPaise": 35000}], "methods": [{"method": "CASH", "amountPaise": 35000}]}, Reason="meant for the bag", Actor=admin)
    db.refresh(fee)
    db.refresh(bagi)
    assert (fee.status, bagi.status) == ("PENDING", "PAID")


def test_edit_cannot_take_back_advance_already_used():
    db, admin, monthly, bag = _setup()
    payment = _pay(db, admin, allocations=[], methods=[("CASH", 110000, None)], keepAdvance=True)
    inv.GenerateInvoices(db, Request=_req([monthly], students=("s1",), month=12), IdempotencyKey="gen-00000004", Actor=admin)
    with pytest.raises(HTTPException) as excinfo:
        rec.EditPayment(db, PaymentId=payment["paymentId"], Request={"methods": [{"method": "CASH", "amountPaise": 50000}], "keepAdvance": True}, Reason="less", Actor=admin)
    assert _code(excinfo) == "ADVANCE_ALREADY_USED"


# --- cancel ----------------------------------------------------------------------

def test_cancel_puts_the_invoices_back_and_keeps_the_receipt():
    db, admin, monthly, bag = _setup()
    fee = _invoice(db, "s1", "Monthly Fee")
    payment = _pay(db, admin, allocations=[(fee.id, 110000, 0)])
    with pytest.raises(HTTPException) as excinfo:
        rec.CancelPayment(db, PaymentId=payment["paymentId"], Reason="", Actor=admin)
    assert _code(excinfo) == "REASON_REQUIRED"
    cancelled = rec.CancelPayment(db, PaymentId=payment["paymentId"], Reason="cheque bounced", Actor=admin)
    assert cancelled["status"] == "CANCELLED" and cancelled["cancelReason"] == "cheque bounced"
    db.refresh(fee)
    assert (fee.status, fee.paid_paise) == ("PENDING", 0)
    with pytest.raises(HTTPException) as excinfo:
        rec.CancelPayment(db, PaymentId=payment["paymentId"], Reason="again", Actor=admin)
    assert _code(excinfo) == "PAYMENT_ALREADY_CANCELLED"
    with pytest.raises(HTTPException) as excinfo:
        rec.EditPayment(db, PaymentId=payment["paymentId"], Request={"methods": [{"method": "CASH", "amountPaise": 1}]}, Reason="x", Actor=admin)
    assert _code(excinfo) == "PAYMENT_CANCELLED"
    # The number is not given back.
    assert db.get(PaymentNumberSequence, "RECEIPT").next_number == 632


def test_cancel_is_refused_while_its_advance_is_used():
    db, admin, monthly, bag = _setup()
    payment = _pay(db, admin, allocations=[], methods=[("CASH", 110000, None)], keepAdvance=True)
    inv.GenerateInvoices(db, Request=_req([monthly], students=("s1",), month=12), IdempotencyKey="gen-00000005", Actor=admin)
    with pytest.raises(HTTPException) as excinfo:
        rec.CancelPayment(db, PaymentId=payment["paymentId"], Reason="x", Actor=admin)
    assert _code(excinfo) == "ADVANCE_ALREADY_USED" and "MP-INV-" in excinfo.value.detail["message"]


# --- account ---------------------------------------------------------------------

def test_student_account_totals_and_statement_agree():
    db, admin, monthly, bag = _setup()
    fee, bagi = _invoice(db, "s1", "Monthly Fee"), _invoice(db, "s1", "MathPath Bag")
    _pay(db, admin, allocations=[(fee.id, 100000, 10000)], methods=[("CASH", 100000, None)], discountReason="concession")
    _pay(db, admin, allocations=[(bagi.id, 20000, 0)], methods=[("UPI", 50000, "UTR5")], keepAdvance=True)
    account = rec.StudentAccount(db, "s1")
    totals = account["totals"]
    assert totals["invoicedDisplay"] == "₹1,450.00"
    assert totals["receivedDisplay"] == "₹1,500.00"
    assert totals["discountDisplay"] == "₹100.00"
    assert totals["dueDisplay"] == "₹150.00"
    assert totals["advanceDisplay"] == "₹300.00"
    assert totals["overdueDisplay"] == "₹150.00"  # due 11 Nov, today 20 Nov
    # Running balance = due - advance = -150 (in credit).
    assert account["statement"][0]["balancePaise"] == totals["duePaise"] - totals["advancePaise"] == -15000
    assert [row["invoiceNumber"] for row in account["unpaidInvoices"]] == [bagi.invoice_number]
    assert len(account["payments"]) == 2 and account["student"]["centreName"] == "Rajarhat"


# --- list, export, PDF -----------------------------------------------------------

def test_list_filters_and_totals_by_method():
    db, admin, monthly, bag = _setup()
    fee1, fee2 = _invoice(db, "s1", "Monthly Fee"), _invoice(db, "s2", "Monthly Fee")
    _pay(db, admin, allocations=[(fee1.id, 110000, 0)], methods=[("CASH", 50000, None), ("UPI", 60000, "UTR-A")])
    second = _pay(db, admin, student="s2", allocations=[(fee2.id, 110000, 0)], methods=[("UPI", 110000, "UTR-B")])
    third = _pay(db, admin, student="s2", allocations=[], methods=[("CASH", 1000, None)], keepAdvance=True)
    rec.CancelPayment(db, PaymentId=third["paymentId"], Reason="test", Actor=admin)
    everything = rec.ListPayments(db, Filters={})
    assert everything["totalCount"] == 3 and everything["totals"]["receivedDisplay"] == "₹2,200.00"
    assert {row["methodLabel"]: row["amountDisplay"] for row in everything["totals"]["byMethod"]} == {"UPI": "₹1,700.00", "Cash": "₹500.00"}
    assert rec.ListPayments(db, Filters={"method": "CASH"})["totalCount"] == 2
    assert rec.ListPayments(db, Filters={"search": "utr-b"})["payments"][0]["receiptNumber"] == second["receiptNumber"]
    assert rec.ListPayments(db, Filters={"status": "CANCELLED"})["totalCount"] == 1
    assert rec.ListPayments(db, Filters={"centreId": "NONE"})["totalCount"] == 2
    book = load_workbook(io.BytesIO(BuildPaymentsWorkbook(rec.PaymentsForExport(db, Filters={}))))
    assert book["Payments"].cell(row=book["Payments"].max_row, column=10).value == 2200.0
    assert [tuple(r) for r in book["By Method"].iter_rows(values_only=True)][-1] == ("Total", 2200.0)


def test_receipt_pdf_has_everything():
    db, admin, monthly, bag = _setup()
    fee = _invoice(db, "s1", "Monthly Fee")
    payment = _pay(db, admin, allocations=[(fee.id, 100000, 10000)], methods=[("CASH", 50000, None), ("UPI", 150000, "UTR-PDF-1")], discountReason="sibling", keepAdvance=True)
    row = db.get(PaymentReceipt, payment["paymentId"])
    content = invoice_pdf.RenderReceiptsPdf([(json.loads(row.snapshot_json), rec.PaymentPayload(db, row, Names={}))])
    text = PdfReader(io.BytesIO(content)).pages[0].extract_text()
    for expected in ["MONEY RECEIPT", "MP-MRCPT-631", "RECEIVED", "Aarav Sen", fee.invoice_number, "Cash", "UPI", "UTR-PDF-1",
                     "2,000.00", "Discount given", "Kept as advance", "1,000.00", "Rupees Two Thousand Only", "sibling", "Authorised Signatory", "Money Receipt"]:
        assert expected in text, expected
    rec.EditPayment(db, PaymentId=row.id, Request={"allocations": [{"invoiceId": fee.id, "amountPaise": 100000, "discountPaise": 10000}], "methods": [{"method": "CASH", "amountPaise": 100000}], "discountReason": "sibling"}, Reason="fix", Actor=admin)
    rec.CancelPayment(db, PaymentId=row.id, Reason="bounced", Actor=admin)
    db.refresh(row)
    text = PdfReader(io.BytesIO(invoice_pdf.RenderReceiptsPdf([(json.loads(row.snapshot_json), rec.PaymentPayload(db, row, Names={}))]))).pages[0].extract_text()
    assert "CANCELLED" in text and "bounced" in text and fee.invoice_number in text


# --- API -------------------------------------------------------------------------

def test_payment_api_end_to_end():
    from app.main import app

    db, admin, monthly, bag = _setup()
    fee = _invoice(db, "s1", "Monthly Fee")
    app.dependency_overrides[get_current_user] = lambda: admin
    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)
    try:
        account = client.get("/api/admin/payments/students/s1/account").json()
        assert account["receiptNumbering"]["nextNumber"] == "MP-MRCPT-631" and len(account["unpaidInvoices"]) == 2
        staff = client.get("/api/admin/payments/staff").json()
        assert staff["currentUserId"] == "admin-1"
        body = {"studentId": "s1", "idempotencyKey": "api-pay-0001", "paymentDate": "2026-10-08", "allocations": [{"invoiceId": fee.id, "amountPaise": 110000}], "methods": [{"method": "UPI", "amountPaise": 110000, "reference": "UTR-API"}]}
        made = client.post("/api/admin/payments/receipts", json=body)
        assert made.status_code == 200 and made.json()["receiptNumber"] == "MP-MRCPT-631"
        payment_id = made.json()["paymentId"]
        assert client.post("/api/admin/payments/receipts", json=body).json()["replayed"] is True
        listed = client.get("/api/admin/payments/receipts", params={"method": "UPI"}).json()
        assert listed["totalCount"] == 1
        assert client.get(f"/api/admin/payments/receipts/{payment_id}").json()["allocations"][0]["invoiceNumber"] == fee.invoice_number
        pdf = client.get(f"/api/admin/payments/receipts/{payment_id}/pdf")
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and "MP-MRCPT-631" in pdf.headers["content-disposition"]
        bulk = client.post("/api/admin/payments/receipts/pdf", json={"filters": {"status": "RECORDED"}})
        assert bulk.status_code == 200 and len(PdfReader(io.BytesIO(bulk.content)).pages) == 1
        xlsx = client.get("/api/admin/payments/receipts/export")
        assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
        edit = client.put(f"/api/admin/payments/receipts/{payment_id}", json={"reason": "fix ref", "allocations": [{"invoiceId": fee.id, "amountPaise": 110000}], "methods": [{"method": "UPI", "amountPaise": 110000, "reference": "UTR-API-2"}]})
        assert edit.status_code == 200 and edit.json()["methods"][0]["reference"] == "UTR-API-2"
        cancel = client.post(f"/api/admin/payments/receipts/{payment_id}/cancel", json={"reason": "test"})
        assert cancel.status_code == 200 and cancel.json()["status"] == "CANCELLED"
        assert client.post("/api/admin/payments/students/s1/apply-advance").status_code == 409
        assert client.get("/api/admin/payments/receipts/nope").status_code == 404
        assert client.get("/api/admin/payments/students/nope/account").status_code == 404
    finally:
        app.dependency_overrides.clear()


def test_advance_plan_is_first_in_first_out():
    assert rec.PlanAdvance([500, 300], [400, 200, 300]) == [(0, 0, 400), (0, 1, 100), (1, 1, 100), (1, 2, 200)]
    assert rec.PlanAdvance([], [100]) == [] and rec.PlanAdvance([100], []) == []
