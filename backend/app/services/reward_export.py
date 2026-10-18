"""Rewards history as an Excel sheet (2026-10-18). Replaces the old portal's
Reward Report export. India time; no phone numbers."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.birthday_service import INDIA

STATUS_LABELS = {"UNLOCKED": "Unlocked", "READY": "Ready", "COLLECTED": "Collected", "CANCELLED": "Cancelled", "REJECTED": "Rejected"}

COLUMNS = [
    ("Unlocked", 18), ("Student", 26), ("Student ID", 14), ("Level", 10), ("Reward", 34), ("Coins", 9), ("Pickup Code", 13),
    ("Status", 12), ("Ready", 18), ("Ready By", 20), ("Collected", 18), ("Handed Over By", 20), ("Closed", 18), ("Closed By", 20), ("Reason", 36),
]


def _When(Value: str | None):
    if not Value:
        return None
    Moment = datetime.fromisoformat(Value).astimezone(INDIA)
    return Moment.replace(tzinfo=None)


def BuildRewardsWorkbook(Orders: list[dict]) -> bytes:
    Book = Workbook()
    Sheet = Book.active
    Sheet.title = "Rewards"
    Sheet.append([Title for Title, _ in COLUMNS])
    Fill = PatternFill("solid", fgColor="E8EAF6")
    for Index, (_, Width) in enumerate(COLUMNS, start=1):
        Cell = Sheet.cell(row=1, column=Index)
        Cell.font = Font(bold=True)
        Cell.fill = Fill
        Cell.alignment = Alignment(vertical="center")
        Sheet.column_dimensions[get_column_letter(Index)].width = Width
    Sheet.freeze_panes = "A2"
    for Order in Orders:
        Student = Order.get("student") or {}
        Sheet.append([
            _When(Order.get("unlockedAt")),
            Student.get("name") or "",
            Student.get("studentCode") or "",
            Student.get("levelCode") or "",
            Order.get("itemName") or "",
            Order.get("price"),
            Order.get("pickupCode") or "",
            STATUS_LABELS.get(Order.get("status"), Order.get("status") or ""),
            _When(Order.get("readyAt")),
            Order.get("readyBy") or "",
            _When(Order.get("collectedAt")),
            Order.get("collectedBy") or "",
            _When(Order.get("closedAt")),
            Order.get("closedBy") or "",
            Order.get("closeReason") or "",
        ])
        Row = Sheet.max_row
        for Column in (1, 9, 11, 13):
            Sheet.cell(row=Row, column=Column).number_format = "DD MMM YYYY, HH:MM"
    Out = BytesIO()
    Book.save(Out)
    return Out.getvalue()
