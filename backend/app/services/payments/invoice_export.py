"""Invoices as an Excel sheet (Payments Phase 2, 2026-10-08).

Amounts are written as numbers in rupees (two decimals) so the sheet can be
summed and filtered; a totals row at the bottom leaves cancelled invoices out,
as the Invoices screen does."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS: list[tuple[str, str, int]] = [
    ("Invoice No.", "invoiceNumber", 16),
    ("Invoice Date", "invoiceDate", 13),
    ("Due Date", "dueDate", 13),
    ("Student", "studentName", 26),
    ("Student ID", "studentCode", 13),
    ("Level", "levelCode", 10),
    ("Centre", "centreName", 14),
    ("Fee", "feeName", 26),
    ("Period", "periodLabel", 15),
    ("Taxable (Rs)", "taxablePaise", 13),
    ("CGST (Rs)", "cgstPaise", 11),
    ("SGST (Rs)", "sgstPaise", 11),
    ("Amount (Rs)", "amountPaise", 13),
    ("Paid (Rs)", "paidPaise", 12),
    ("Balance (Rs)", "balancePaise", 13),
    ("Status", "statusLabel", 12),
    ("Overdue", "isOverdue", 9),
    ("Cancel Reason", "cancelReason", 28),
]
MONEY_KEYS = {"taxablePaise", "cgstPaise", "sgstPaise", "amountPaise", "paidPaise", "balancePaise"}


def BuildInvoicesWorkbook(Rows: list[dict[str, Any]]) -> bytes:
    Book = Workbook()
    Sheet = Book.active
    Sheet.title = "Invoices"
    Sheet.append([Title for Title, _, _ in COLUMNS])
    HeaderFill = PatternFill("solid", fgColor="E8EAF6")
    for Index, (_, _, Width) in enumerate(COLUMNS, start=1):
        Cell = Sheet.cell(row=1, column=Index)
        Cell.font = Font(bold=True)
        Cell.fill = HeaderFill
        Cell.alignment = Alignment(vertical="center")
        Sheet.column_dimensions[get_column_letter(Index)].width = Width
    Sheet.freeze_panes = "A2"

    Totals = {Key: 0 for Key in MONEY_KEYS}
    for Row in Rows:
        Values = []
        for _, Key, _ in COLUMNS:
            Value = Row.get(Key)
            if Key in MONEY_KEYS:
                Value = round(int(Value or 0) / 100, 2)
            elif Key == "isOverdue":
                Value = "Yes" if Value else ""
            Values.append(Value if Value is not None else "")
        Sheet.append(Values)
        if Row.get("status") != "CANCELLED":
            for Key in MONEY_KEYS:
                Totals[Key] += int(Row.get(Key) or 0)

    for Index, (_, Key, _) in enumerate(COLUMNS, start=1):
        if Key in MONEY_KEYS:
            for RowIndex in range(2, Sheet.max_row + 1):
                Sheet.cell(row=RowIndex, column=Index).number_format = "#,##0.00"

    if Rows:
        Sheet.append([])
        TotalRow = ["Total (excluding cancelled)"] + [""] * (len(COLUMNS) - 1)
        for Index, (_, Key, _) in enumerate(COLUMNS):
            if Key in MONEY_KEYS:
                TotalRow[Index] = round(Totals[Key] / 100, 2)
        Sheet.append(TotalRow)
        Last = Sheet.max_row
        for Index, (_, Key, _) in enumerate(COLUMNS, start=1):
            Cell = Sheet.cell(row=Last, column=Index)
            Cell.font = Font(bold=True)
            if Key in MONEY_KEYS:
                Cell.number_format = "#,##0.00"

    Buffer = BytesIO()
    Book.save(Buffer)
    return Buffer.getvalue()
