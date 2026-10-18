"""Birthdays as an Excel sheet (2026-10-18). No phone numbers."""
from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STATUS_LABELS = {
    "SEEN": "Seen",
    "SENT": "Sent",
    "WAITING": "Waiting",
    "SCHEDULED": "Scheduled",
    "NOT_SENT": "Not Sent",
}

COLUMNS = [("Date", 12), ("Day", 11), ("Student", 26), ("Student ID", 14), ("Turns", 8), ("Level", 10), ("Teacher", 22), ("Centre", 14), ("Wish", 12)]


def BuildBirthdaysWorkbook(Entries: list[dict]) -> bytes:
    Book = Workbook()
    Sheet = Book.active
    Sheet.title = "Birthdays"
    Sheet.append([Title for Title, _ in COLUMNS])
    Fill = PatternFill("solid", fgColor="E8EAF6")
    for Index, (_, Width) in enumerate(COLUMNS, start=1):
        Cell = Sheet.cell(row=1, column=Index)
        Cell.font = Font(bold=True)
        Cell.fill = Fill
        Cell.alignment = Alignment(vertical="center")
        Sheet.column_dimensions[get_column_letter(Index)].width = Width
    Sheet.freeze_panes = "A2"
    for Entry in Entries:
        Day = date.fromisoformat(Entry["date"])
        Sheet.append([
            Day,
            Day.strftime("%A"),
            Entry.get("name") or "",
            Entry.get("studentCode") or "",
            Entry.get("age"),
            Entry.get("levelCode") or "",
            Entry.get("teacherName") or "",
            Entry.get("centreName") or "",
            STATUS_LABELS.get(Entry.get("status"), ""),
        ])
        Sheet.cell(row=Sheet.max_row, column=1).number_format = "DD MMM YYYY"
    Out = BytesIO()
    Book.save(Out)
    return Out.getvalue()
